"""管理者向け手順書「デフォルト設定の編集ガイド」（PowerPoint）を作る。

使い方：python3 build_deck.py <出力 .pptx> [参照デッキ .pptx]
- 参照デッキ：技術レポートPowerPoint（reference.pptx）。スライドマスター・テーマ・フォントをそのまま使い、
  元のスライドは消して、同じ色・位置・文字の大きさで新しいスライドを置く。
- 見た目の部品は、解説資料（docs/kernel_guide/_src/pptkit.py）と共通のものを使う。
- 表の初期値は matlab_app/+alignest/defaultSettings.m・recommendedMethods.m と画面の定数に合わせて書いている。
  初期値を変えたときは、この表も直して作り直す。
出力は別名（.building.pptx）に書いてから置き換える。出力先に PowerPoint のロックファイル（~$…）があれば止める。
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

from pptx import Presentation

HERE = Path(__file__).resolve().parent
# 参照デッキの見た目の部品（解説資料と共通。同じ処理を2か所に書かないため）
sys.path.insert(0, str(HERE.parent.parent / "kernel_guide" / "_src"))
from pptkit import (BODY, CORAL, CORAL_TINT, GOLD, GOLD_TINT, INK, LIGHT, LIGHT2, LINE, MUTED, NAVY, NAVY_PANEL,  # noqa: E402
                    PANEL, PURPLE, PURPLE_TINT, TEAL, TEAL_TINT, WHITE, YELLOW, background, box, chrome, circle, line,
                    notes, point_bar, table, text)

DEFAULT_REFERENCE = Path.home() / ".claude/skills/artifact-template-powerpoint-2/assets/reference.pptx"
DECK_NAME = "デフォルト設定の編集ガイド"
DECK_TITLE = "デフォルト設定の編集ガイド（管理者向け）"
CODE_FONT = "Courier New"
DEFAULTS_FILE = "alignment_estimation_defaults.json"
TEST_COUNT = 205   # runAllTests の件数（2026-10-31 時点。テストを足したら直す）


class Deck:
    def __init__(self, reference: Path):
        self.prs = Presentation(str(reference))
        self.layout = self.prs.slides[0].slide_layout
        # 参照デッキのスライドを外す（マスター・テーマ・フォントは残る）
        id_list = self.prs.slides._sldIdLst
        for slide_id in list(id_list):
            self.prs.part.drop_rel(slide_id.rId)
            id_list.remove(slide_id)
        self.page = 0

    def slide(self, dark=False):
        slide = self.prs.slides.add_slide(self.layout)
        for placeholder in list(slide.placeholders):
            placeholder._element.getparent().remove(placeholder._element)
        background(slide, NAVY if dark else WHITE)
        self.page += 1
        return slide

    def body(self, chapter, title):
        slide = self.slide()
        chrome(slide, chapter, title, self.page, DECK_NAME)
        return slide


def steps(slide, x, y, w, items, color=TEAL, gap=0.78, size=14):
    """番号の丸と、太字の見出し＋説明の段"""
    for index, (heading, body) in enumerate(items):
        yy = y + index * gap
        circle(slide, x, yy + 0.05, 0.48, color, str(index + 1), size=15)
        text(slide, x + 0.65, yy, w - 0.65, 0.34, heading, size=size, color=INK, bold=True)
        text(slide, x + 0.65, yy + 0.33, w - 0.65, gap - 0.33, body, size=12, color=BODY)


def code_box(slide, x, y, w, h, title, lines, color=NAVY, size=11):
    """題名つきのコードの箱（等幅の文字）"""
    box(slide, x, y, w, 0.42, fill=color)
    text(slide, x + 0.15, y, w - 0.3, 0.42, title, size=13, color=WHITE, bold=True, anchor="m")
    box(slide, x, y + 0.42, w, h - 0.42, fill=PANEL, line=LINE)
    text(slide, x + 0.15, y + 0.52, w - 0.3, h - 0.6, lines, size=size, color=INK, font=CODE_FONT)


def settings_table(slide, rows, y=1.5, widths=(3.4, 1.85, 5.08, 1.8), size=11, row_h=0.36):
    table(slide, 0.6, y, list(widths), [["項目（設定ファイルの名前）", "初期値", "意味", "画面"]] + rows, size=size,
          header_size=12, row_h=row_h, aligns=["l", "l", "l", "l"])


def build(output: Path, reference: Path):
    deck = Deck(reference)

    # 1. 表紙
    s = deck.slide(dark=True)
    text(s, 0.8, 0.9, 7.0, 0.4, "管理者向け手順書 ｜ 面内推定評価ツール（MATLAB）", size=14, color=LIGHT)
    text(s, 0.8, 1.9, 7.2, 1.0, "デフォルト設定の編集ガイド", size=34, color=WHITE, bold=True)
    text(s, 0.8, 3.05, 6.8, 1.3, "評価の条件・比べる手法・画面の初期値を\nどこで、どう変えるか", size=20, color=LIGHT2)
    text(s, 0.8, 6.35, 6.8, 0.35, "2026年10月 ｜ MATLAB R2022b・matlab_app フォルダ", size=12, color=LIGHT)
    box(s, 8.2, 0.95, 4.53, 5.6, fill=NAVY_PANEL)
    text(s, 8.5, 1.15, 4.0, 0.4, "変え方は3つ", size=16, color=WHITE, bold=True)
    ways = [(TEAL, "A", "設定ファイルを配る", "利用者が「設定を読み込む…」で読む"),
            (CORAL, "B", "既定の設定ファイルを置く", "起動したときの初期設定になる（exe の作り直し不要）"),
            (PURPLE, "C", "コードの初期値を変える", "defaultSettings.m を直して exe を作り直す")]
    for index, (color, label, heading, body) in enumerate(ways):
        y = 1.8 + index * 1.55
        circle(s, 8.5, y, 0.62, color, label, size=18)
        text(s, 9.3, y - 0.05, 3.3, 0.4, heading, size=15, color=WHITE, bold=True)
        text(s, 9.3, y + 0.38, 3.3, 0.9, body, size=12.5, color=LIGHT2)
    notes(s, "管理者が、面内推定評価ツールの初期設定（評価の条件・比べる手法・画面の初期値）を変えるときの手順書。"
             "3つの変え方（設定ファイルを配る・既定の設定ファイルを置く・コードの初期値を変える）と、項目ごとの初期値、注意点、確かめ方をまとめる。")

    # 2. この資料で分かること
    s = deck.body("", "この資料で分かること")
    box(s, 0.6, 1.5, 5.6, 2.35, fill=PANEL)
    text(s, 0.85, 1.65, 5.2, 0.4, "対象者と前提", size=16, color=INK, bold=True)
    text(s, 0.85, 2.1, 5.2, 1.7, ["アプリを部署で保守・配布する管理者",
                                  "MATLAB R2022b と matlab_app フォルダ（exe を作るときは MATLAB Compiler）",
                                  "手順 A・B は MATLAB のコードを書かずにできる",
                                  "手順 C は MATLAB のファイルを開いて値を書き換える"], size=13, color=BODY, bullets=True,
         space_after=4)
    box(s, 0.6, 4.05, 5.6, 2.75, fill=TEAL_TINT)
    text(s, 0.85, 4.2, 5.2, 0.4, "この資料で扱わないこと", size=16, color=TEAL, bold=True)
    text(s, 0.85, 4.65, 5.2, 2.1, ["推定の手法の中身を変えること（計算の関数の修正）",
                                   "入力の列の名前が変わったときの対応（管理者向け資料 HTML）",
                                   "exe を配る手順の詳細（Windows 手順書）"], size=13, color=BODY, bullets=True, space_after=4)
    chapters = [("01", "設定の仕組み", "初期値に、ファイルと画面の変更を重ねる流れ", TEAL),
                ("02", "3つの変え方と手順", "設定ファイル・既定の設定ファイル・コード", CORAL),
                ("03", "設定の項目と初期値", "data・factors・sampling・evaluation・overlay・methods・highOrder", PURPLE),
                ("04", "編集の例・注意・確かめ方", "JSON とコードの例、型・単位、テスト", GOLD)]
    for index, (number, heading, body, color) in enumerate(chapters):
        y = 1.5 + index * 1.33
        box(s, 6.6, y, 6.13, 1.15, fill=WHITE, line=LINE)
        circle(s, 6.85, y + 0.27, 0.62, color, number, size=15, font="Times New Roman")
        text(s, 7.7, y + 0.15, 4.9, 0.42, heading, size=17, color=INK, bold=True)
        text(s, 7.7, y + 0.6, 4.9, 0.45, body, size=13, color=BODY)
    notes(s, "前半で仕組みと3つの変え方を説明し、後半で項目ごとの初期値の一覧と、編集の例・注意点・確かめ方をまとめる。")

    # 3. 設定の流れ
    s = deck.body("01　設定の仕組み", "設定の流れ ― 初期値に、ファイルと画面の変更を重ねる")
    flow = [(PURPLE, "コードの初期値", "defaultSettings.m\nrecommendedMethods.m", "C"),
            (CORAL, "既定の設定ファイル", DEFAULTS_FILE.replace("estimation_", "estimation_\n") + "\n（あれば起動時に重ねる）",
             "B"),
            (TEAL, "画面で変える", "「2. 評価の設定」\nの欄", ""),
            (TEAL, "設定ファイル", "「設定を保存…」\n「設定を読み込む…」", "A"),
            (NAVY, "評価結果の Excel", "「設定」シートに\n使った設定を残す", "")]
    width = 2.15
    for index, (color, heading, body, label) in enumerate(flow):
        x = 0.6 + index * 2.5
        box(s, x, 1.75, width, 2.2, fill=WHITE, line=color, line_width=1.5)
        box(s, x, 1.75, width, 0.5, fill=color)
        text(s, x, 1.75, width, 0.5, heading, size=14, color=WHITE, bold=True, align="c", anchor="m")
        text(s, x + 0.1, 2.35, width - 0.2, 1.5, body, size=12, color=BODY, align="c", anchor="m")
        if label:
            circle(s, x + width / 2 - 0.25, 4.1, 0.5, color, label, size=14)
        if index < len(flow) - 1:
            line(s, x + width + 0.05, 2.85, x + 2.45, 2.85, color=MUTED, width=1.75)
    text(s, 0.6, 4.75, 12.13, 0.9, ["読み込むときは、ファイルに書いてある項目だけを初期値に重ねる（alignest.loadSettings）。"
                                    "書いていない項目・以前の版に無かった項目は初期値になる。",
                                    "以前の版の「Wafer 半径に対する比」の項目（radialZoneEdges・outerRadiusRatio）は、"
                                    "読み込むときに 比 × 150 mm の mm の項目に直す。"], size=13, color=BODY, bullets=True,
         space_after=4)
    point_bar(s, 5.85, "だから、設定ファイル・既定の設定ファイルには「変えたい項目だけ」を書けばよい（ほかは初期値のまま）",
              icon="i", icon_color=BODY)
    notes(s, "設定は、コードの初期値 → 既定の設定ファイル（あれば）→ 画面の操作 の順に重なる。"
             "設定を保存すると全項目が JSON に書かれ、評価結果の Excel にも同じ JSON が残る。")

    # 4. 3つの変え方の比較
    s = deck.body("01　設定の仕組み", "3つの変え方 ― exe を作り直すかどうかで選ぶ")
    rows = [["変え方", "向いている場面", "exe の作り直し", "変わる範囲", "元に戻すには"],
            ["A 設定ファイルを配る", "評価の目的ごとに条件を変える\n人によって条件を変える", "不要",
             "読み込んだ人の、そのときの設定", "「初期設定に戻す」"],
            ["B 既定の設定ファイルを置く", "部署の標準の初期設定をそろえる\n（全員が同じ初期値から始める）", "不要",
             "そのフォルダから起動した人の初期設定", "ファイルを消す・名前を変える"],
            ["C コードの初期値を変える", "アプリの標準そのものを変える\n（新しい項目の初期値を決める）",
             "必要（MATLAB Compiler）", "アプリを使う全員", "Git で前の版に戻す"]]
    table(s, 0.6, 1.55, [2.9, 3.6, 1.65, 2.3, 1.68], rows, size=12, row_h=0.95, bold_first=True,
          fills=[None, TEAL_TINT, CORAL_TINT, PURPLE_TINT])
    point_bar(s, 5.85, "まず A・B（exe の作り直し不要）を使う。全員の標準を恒久的に変えるときだけ C にする",
              icon="✓", icon_color=TEAL)
    notes(s, "A と B はファイルを置くだけなので、会社の PC でも MATLAB Compiler なしでできる。"
             "C はコードの初期値を変えるので、テストと exe の作り直しが必要。")

    # 5. 手順A
    s = deck.body("02　3つの変え方と手順", "手順 A：設定ファイルを作って配る")
    steps(s, 0.6, 1.55, 6.4, [
        ("「2. 評価の設定」で条件を決める", "学習 Shot の選び方と数・制約・比べる手法・評価の条件・Overlay 補正を選ぶ"),
        ("「設定を保存…」で JSON に保存する", "名前の例：standard_settings.json（UTF-8 の文字ファイル）"),
        ("共有フォルダに置き、使い方を伝える", "どの評価に使う設定か、ファイル名とあわせて伝える"),
        ("利用者は「設定を読み込む…」で読む", "データを読み込んだ後でもよい。読み込むと画面の欄がその値になる"),
        ("戻すときは「初期設定に戻す」", "アプリの初期設定（B があればその値）に戻る")], color=TEAL, gap=0.88)
    box(s, 7.35, 1.55, 5.38, 2.55, fill=TEAL_TINT)
    text(s, 7.6, 1.7, 4.9, 0.4, "設定ファイルに入るもの", size=15, color=TEAL, bold=True)
    text(s, 7.6, 2.15, 4.9, 1.9, ["データの条件（有効半径・座標の単位・行の除外・Scan 方向の付け方）",
                                  "質的変数、学習 Shot の選び方・数・制約・必ず計測する Shot",
                                  "Excel ファイルで指定した Shot・Mark（パターン・優先順）、Mark 単位の計画",
                                  "比べる手法、評価の条件、Overlay 補正のモデルと式"], size=12, color=BODY, bullets=True,
         space_after=3)
    box(s, 7.35, 4.3, 5.38, 1.45, fill=PANEL)
    text(s, 7.6, 4.45, 4.9, 0.4, "入らないもの", size=15, color=INK, bold=True)
    text(s, 7.6, 4.9, 4.9, 0.8, ["計測ファイル・列の対応（読み込むときに選ぶ）",
                                 "画面の色（ライト・ダーク。人ごとに保存される）"], size=12, color=BODY, bullets=True)
    point_bar(s, 6.0, "評価結果の Excel の「設定」シートにも同じ JSON が残る。過去の評価と同じ条件にしたいときに使える",
              icon="i", icon_color=BODY, h=0.8, size=14)
    notes(s, "設定ファイルは画面の「設定を保存…」で作る。手で書いた JSON でもよい（変えたい項目だけでよい）。")

    # 6. 手順B
    s = deck.body("02　3つの変え方と手順", "手順 B：既定の設定ファイルを置く ― 全員の初期値を変える")
    steps(s, 0.6, 1.55, 6.3, [
        ("JSON を作る", "手順 A の「設定を保存…」で作るか、変えたい項目だけを書く（右の例）"),
        ("名前を " + DEFAULTS_FILE + " にする", "この名前のファイルだけを読む"),
        ("起動するフォルダに置く", "MATLAB：現在のフォルダ（matlab_app）。exe：ショートカットの作業フォルダー（ふつう exe のフォルダ）"),
        ("起動して確かめる", "「3. 実行」の記録に「既定の設定ファイルを読み込みました」と出る。「初期設定に戻す」もこの値"),
        ("読めないときは初期設定で開く", "警告と理由が出る。ファイルを直すか、名前を変えて開き直す")], color=CORAL, gap=0.88)
    code_box(s, 7.25, 1.55, 5.48, 3.4, DEFAULTS_FILE + "（例）", [
        "{",
        '  "sampling": {',
        '    "method": ["ioptimal", "doptimal"],',
        '    "nShotsStart": 10, "nShotsEnd": 50,',
        '    "nShotsStep": 10',
        "  },",
        '  "evaluation": {',
        '    "outerRadiusMm": 130,',
        '    "outlierMethod": "mad"',
        "  }",
        "}"], color=CORAL, size=11.5)
    text(s, 7.25, 5.05, 5.48, 0.75, "書いた項目だけが変わる（ほかは初期値）。学習 Shot 数は開始・終了・刻みだけ書けば、並びも合わせる。",
         size=12, color=BODY)
    point_bar(s, 6.0, "exe を作り直さずに、部署の全員の初期設定をそろえられる（ファイルを消せば元に戻る）",
              icon="✓", icon_color=CORAL, h=0.8, size=14)
    notes(s, "既定の設定ファイルは alignest.siteDefaults が起動時に読む。今のフォルダ（pwd）にある " + DEFAULTS_FILE +
          " だけが対象。読めないときはアプリの初期設定で開き、警告を出す。")

    # 7. 手順C
    s = deck.body("02　3つの変え方と手順", "手順 C：コードの初期値を変える")
    steps(s, 0.6, 1.5, 7.0, [
        ("元の状態を残す", "Git のブランチを作るか、matlab_app フォルダを複製する"),
        ("初期値を書き換える", "defaultSettings.m・recommendedMethods.m。型・単位・項目の名前を保つ"),
        ("テストを流す", "cd tests → runAllTests。初期値を確かめるテストの期待値も見直す"),
        ("画面で確かめる", "「初期設定に戻す」→ 新しい値。「設定を保存…」の JSON も見る"),
        ("文書を直す", "README・Windows 手順書・この資料の初期値の表"),
        ("exe を作り直して配る", "buildStandaloneApp(makeInstaller=true)（MATLAB Compiler の入った PC）")],
        color=PURPLE, gap=0.8, size=13.5)
    box(s, 7.95, 1.5, 4.78, 4.75, fill=PURPLE_TINT)
    text(s, 8.2, 1.65, 4.3, 0.4, "直すファイル", size=15, color=PURPLE, bold=True)
    files = [("defaultSettings.m", "data・factors・sampling・evaluation・overlay・highOrder の初期値"),
             ("recommendedMethods.m", "初期設定（GPR だけ）・推奨の組み合わせ・すべて・なし と手法の細かい設定（options）"),
             ("+appui/Core.m", "画面の定数（先頭の properties (Constant)）"),
             ("estimateRunSeconds.m", "時間の目安（1条件の秒数）"),
             ("exchangeSecondsGuide.m", "交換法の時間の目安")]
    for index, (name, body) in enumerate(files):
        y = 2.15 + index * 0.8
        text(s, 8.2, y, 4.3, 0.34, name, size=12.5, color=INK, bold=True, font=CODE_FONT)
        text(s, 8.2, y + 0.33, 4.3, 0.45, body, size=11.5, color=BODY)
    point_bar(s, 6.45, "C の後は必ずテストと exe の作り直し。過去の評価と比べるときは、前の設定ファイルで評価し直す",
              icon="!", icon_color=PURPLE, h=0.5, size=13)
    notes(s, "コードを変えると、アプリを使う全員の初期値が変わる。変える前の状態を Git で残し、テストを全件流してから配る。")

    # 8. どこに何があるか
    s = deck.body("03　設定の項目と初期値", "どこに何があるか ― 初期値と定数のファイル")
    rows = [["ファイル（matlab_app の中）", "中身", "例"],
            ["+alignest/defaultSettings.m", "評価の条件の初期値（data・factors・sampling・evaluation・overlay・highOrder）",
             "外周とみなす半径 120 mm"],
            ["+alignest/recommendedMethods.m", "比べる手法（初期設定・推奨の組み合わせ・すべて・なし）と手法の細かい設定 options",
             "GPR matern52 × 質的変数4通り"],
            ["+appui/Core.m（先頭）", "画面の大きさ・文字の大きさ・並べる図の数・色", "FONT_SIZE = 16"],
            ["+alignest/estimateRunSeconds.m", "時間の目安（1条件 2 秒、ν を指定した Matérn は 3 倍）", "SECONDS_PER_RUN = 2"],
            ["+alignest/exchangeSecondsGuide.m", "交換法の時間の目安（初期配置1つ × 学習 Shot 1つ）", "0.01 秒"],
            ["+alignest/siteDefaults.m", "既定の設定ファイル（手順 B）を読む", DEFAULTS_FILE],
            ["help/help.html", "「8. ヘルプ」の説明（用語・カーネル・補正の式）", "手法を足したら説明も足す"]]
    table(s, 0.6, 1.5, [3.6, 5.2, 3.33], rows, size=12, row_h=0.56, bold_first=True)
    point_bar(s, 6.2, "設定ファイルの項目の名前は、defaultSettings.m の settings.〈グループ〉.〈項目〉 と同じ",
              icon="i", icon_color=BODY, h=0.6, size=13.5)
    notes(s, "初期値はこの2つの関数にまとまっている。設定ファイル（JSON）の項目名は、関数の中の名前と同じ。")

    # 9. data・factors
    s = deck.body("03　設定の項目と初期値", "設定の項目①：data（データの読み込み）・factors（質的変数）")
    settings_table(s, [
        ["data.waferRadiusMm・validRadiusMm", "150・150", "座標を割る Wafer 半径・この半径 [mm] 以上の点は除く", "1. 有効半径"],
        ["data.coordinateUnits", "すべて mm", "座標の列ごとの入力の単位（mm・um・nm・m）", "1. 座標の単位"],
        ["data.excludeZeroMeasurements・…Unmeasured…", "false・false", "測定値がともに 0・Sts が −1 の行を除く", "1. 行の除外"],
        ["data.directionRule・invertDirection", "\"serpentine\"・false", "Direction が無い行に付ける Scan 方向（一筆書き）・1 と 2 の入れ替え",
         "1. Scan 方向"],
        ["data.coordinateMode", "\"shotMark\"", "座標の定義（shotMark：Shot 中心＋Shot 内、absolute：絶対座標だけ）", "1. 座標の定義"],
        ["data.valueColumn・valueUnit", "\"\"・\"nm\"", "推定する値（空 = X・Y 成分、列の名前 = 膜厚などの1変数）と単位", "1. 推定する値"],
        ["factors.scan・mark・shotRowParity", "true・true・false", "質的変数（Scan 方向・Mark 番号・Shot 行の奇偶）", "2. 質的変数"],
        ["factors.evaluateSingles", "false", "チェックした質的変数を1つだけ使う手法も足して比べる", "2. 1つだけ使う"],
        ["factors.columns", "[]（なし）", "質的変数に使うファイルのその他の列の名前", "2. その他の列"],
        ["factors.separateY ・ factors.y", "false", "Y 成分の質的変数を別に選ぶ（y に同じ形の項目）", "2. Y 成分は別に"]],
        widths=(3.75, 1.9, 4.78, 1.7), row_h=0.4)
    point_bar(s, 6.3, "Wafer 半径を変えるときは validRadiusMm・同心円の区切り・外周の半径（mm）も一緒に見直す",
              icon="!", icon_color=CORAL, h=0.55, size=13)
    notes(s, "data の項目は読み込みの条件。factors は推定に使う質的変数（GPR・GAM・Backfitting）。ANOVA は選ばなくても主な変数を調べる。")

    # 10. sampling（選び方・数）
    s = deck.body("03　設定の項目と初期値", "設定の項目②：sampling（学習 Shot の選び方・数）")
    settings_table(s, [
        ["sampling.method", "\"ioptimal\"", "選び方（複数で比べる）：ioptimal・doptimal・random・excel・mark・markd", "2. 選び方"],
        ["sampling.nShotsStart・End・Step", "10・40・5", "学習 Shot 数の開始・終了・刻み", "2. Shot 数"],
        ["sampling.nShots", "[10 15 … 40]", "上から作った並び（計算に使う。開始〜刻みと合わせる）", "（自動）"],
        ["sampling.countBy", "\"shots\"", "Shot 数で決める（shots）か計測率で決める（rate）か", "2. 学習 Shot 数"],
        ["sampling.rateStart・End・Step", "10・40・10", "計測率 [%] の開始・終了・刻み（countBy = rate）", "2. 計測率"],
        ["sampling.randomRepeats ・ seed", "10 ・ 20261017", "ランダムの繰り返し回数・乱数の種", "2. ランダム"],
        ["sampling.layoutPlans・layoutFile", "[]・\"\"", "Excel ファイルで指定（1行 =「種類|名前|shotNo …」。読み込みで入る）", "2. Excel で指定"],
        ["sampling.requiredShots・requireOuterShots", "[]・false", "必ず計測する shotNo（拡張 I 最適計画）・最外周の Shot をすべて",
         "2. 必ず計測"],
        ["sampling.markPerShot", "1", "Mark 単位：1 Shot あたりの計測 Mark 数（countBy = rate なら計測率）", "2. Mark 単位"],
        ["sampling.markRequirePerShot・markMinPerShot", "true・1", "Mark 単位：Shot ごとに n 点の計測を必須にする・n", "2. Mark 単位"]],
        widths=(3.85, 1.6, 4.98, 1.7), row_h=0.42)
    point_bar(s, 6.3, "method を2つ以上にすると、結果一覧の「計測配置の比較」で選び方どうしを比べられる",
              icon="i", icon_color=BODY, h=0.55, size=13)
    notes(s, "学習 Shot 数は開始・終了・刻みと並び nShots の両方を持つ。コードで変えるときは shotCountList で並びも作る。")

    # 11. sampling（制約・交換法）
    s = deck.body("03　設定の項目と初期値", "設定の項目③：sampling（制約・交換法）")
    settings_table(s, [
        ["sampling.designOrder", "5", "最適計画の基準モデル（XY 多項式の次数。5次で 21 項）", "2. 基準モデル"],
        ["sampling.exchangeStarts", "10", "交換法の初期配置の数（時間はほぼ比例）", "2. 初期配置の数"],
        ["sampling.exchangeMaxIterations", "300", "交換法の反復の上限", "（設定ファイルのみ）"],
        ["constraints.radialZoneEdgesMm", "[67.5 105]", "同心円の区切り：Shot 中心の半径 [mm]", "2. 同心円の区切り"],
        ["constraints.quadrants", "true", "4象限の Shot 数をそろえる", "2. 4象限"],
        ["constraints.scanDirection", "true", "Scan 方向ごとの Shot 数をそろえる", "2. Scan 方向"],
        ["constraints.slack", "1", "許容幅（各区画の数の ±）", "2. 許容幅"],
        ["constraints.markRatio", "true", "Mark 単位：Mark 番号ごとの数を全計測の割合にそろえる", "2. Mark 単位"]],
        row_h=0.44)
    point_bar(s, 6.2, "初期配置の数の目安：1つあたり 約 0.01 秒 × 学習 Shot 数（10 → 30 で約 3 倍）。10〜20 を超えても良さはほぼ同じ",
              icon="i", icon_color=BODY, h=0.6, size=13)
    notes(s, "constraints は sampling の中のグループ（sampling.constraints.〈項目〉）。同心円の区切りは mm の絶対値で、半径 150 mm の Wafer が前提。")

    # 12. evaluation・overlay
    s = deck.body("03　設定の項目と初期値", "設定の項目④：evaluation・overlay（補正）・display（表示）")
    settings_table(s, [
        ["evaluation.outerRadiusMm", "120", "外周とみなす Mark の半径 [mm]（外周の RMSE の境目）", "2. 外周の半径"],
        ["evaluation.removeLinear", "false", "XY 1次を除いてから推定する", "2. 評価"],
        ["evaluation.recordTiming", "false", "条件ごとの処理時間を記録する", "2. 評価"],
        ["evaluation.excludeOutliers", "false", "学習データの異常値を除く", "2. 異常値"],
        ["evaluation.outlierSigma ・ outlierMethod", "3 ・ \"sigma\"", "異常値の倍数 n と判定（sigma：平均と σ、mad：中央値と MAD）",
         "2. 異常値"],
        ["overlay.waferOrder ・ waferUseFormula", "1 ・ false", "Wafer 成分の次数（0 = なし）・式を使うか", "2. Wafer 成分"],
        ["overlay.shotCommon", "\"linear\"", "Shot 共通成分（none・linear・asml・canon・nikon・formula）", "2. Shot 共通成分"],
        ["overlay.perShot", "\"none\"", "Shot ごとの成分（none・translation・linear・formula）", "2. Shot ごとの成分"],
        ["overlay.waferFormula・shotFormula・perShotFormula", "k1 + k3*x + k5*y\nk2 + k4*y + k6*x",
         "式を入力するときの dx・dy の式", "2. dx・dy"],
        ["display.scan1Up・scan1Color・scan2Color", "true・緑・紫", "表示：Scan 方向 1 を上向き（▲）にするか・1・2 の色（RGB 0〜1）",
         "メニュー「表示」"]],
        widths=(3.85, 1.95, 4.68, 1.65), row_h=0.41)
    point_bar(s, 6.3, "Overlay 補正をすべて「なし」（waferOrder 0・shotCommon none・perShot none）にすると補正の評価をしない",
              icon="i", icon_color=BODY, h=0.55, size=13)
    notes(s, "補正の式の書き方と、すべての補正の式は「8. ヘルプ」にまとめている。")

    # 13. methods
    s = deck.body("03　設定の項目と初期値", "設定の項目⑤：methods（比べる手法）")
    settings_table(s, [
        ["methods.featureSets", "\"XYR\"", "説明変数（XY：X,Y、XYR：X,Y,半径）", "2. 説明変数"],
        ["methods.gpr.enabled ・ kernels", "true ・ \"matern52\"", "GPR とカーネル（fitrgp の10種類、matern・ardmatern、periodic・ardperiodic）", "2. GPR"],
        ["methods.gpr.factorModes", "none・basis・additive・mixed", "質的変数の扱い（8通りから）", "2. GPR"],
        ["methods.gpr.maternNu", "2", "matern・ardmatern の ν（複数なら ν ごと）", "2. ν"],
        ["methods.tps・rbf", "false ・ false", "TPS・RBF（factorModes：none・backfit）", "2. TPS・RBF"],
        ["methods.kriging", "通常 false・Universal false", "Kriging（variograms：\"exponential\"）", "2. Kriging"],
        ["methods.gam", "false（スプライン）", "GAM（shapes・gpKernels・gpMaternNu）", "2. GAM"],
        ["methods.polynomial ・ rt", "false（5次）・false（K=5）", "XY 多項式の次数（1〜8）・RT モデルの次数", "2. 多項式・RT"],
        ["methods.zernike", "false（次数 4）", "Fringe Zernike（orders：次数 k は 1〜7。factorModes：none・backfit）", "2. Fringe Zernike"]],
        widths=(3.3, 2.55, 4.6, 1.68), row_h=0.4)
    point_bar(s, 6.2, "初期値は GPR だけ（recommendedMethods.m の \"default\" の部分）。「推奨の組み合わせ」は \"recommended\"、"
              "「すべて選ぶ」は \"all\" の部分",
              icon="i", icon_color=BODY, h=0.6, size=13)
    notes(s, "初期設定（default）：GPR matern52（X,Y,半径）× 質的変数（使わない・同時推定・加法カーネル・混合）の 4 手法だけ。"
             "推奨の組み合わせ（recommended）：それに TPS、Kriging exponential、GAM、XY 5次多項式、RT モデル K=5 を足した 9 手法。"
             "RBFe は廃止（以前の設定ファイルの methods.rbfe は読み飛ばす）。")

    # 14. highOrder（Shot 高次補正）
    s = deck.body("03　設定の項目と初期値", "設定の項目⑥：highOrder（7. Shot高次補正）")
    settings_table(s, [
        ["highOrder.maker", "\"nikon\"", "補正の式（asml・canon・nikon：ルールの式、custom：k1〜k38）", "7. 補正の式"],
        ["highOrder.unit", "\"common\"", "k のくくり（common：共通・scan：Scan 方向・row：行の偶奇・shot）", "7. 補正のくくり"],
        ["highOrder.waferOrder", "1", "Wafer 成分の次数（0〜5、0 = なし。shot は先に求めて引く）", "7. Wafer 成分"],
        ["highOrder.excludedTerms", "[]（すべて使う）", "使わない k の番号（例：[55, 41, 29]。k1〜k6 は外さない）", "7. 使う項と VIF"],
        ["highOrder.outerRadiusMm", "120", "外周の Shot（Shot 中心の半径 [mm] がこの値以上）", "7. 外周の半径"],
        ["highOrder.outerMaker ・ outerName", "\"none\" ・ \"\"", "外周の式（none：分けない）・選んだ組み合わせの名前", "7. 外周の式"],
        ["highOrder.vifThreshold", "10", "「VIF で選ぶ」のしきい値（よく使う値は 5〜10）", "7. VIF のしきい値"],
        ["highOrder.loss ・ huberTuning", "\"huber\" ・ 1.345", "損失関数（ols・huber・lad）・Huber の定数 c", "7. 回帰・Huber"],
        ["highOrder.l2 ・ lambda", "true ・ 0.1", "高次項（2次以上）だけに L2 正則化をかけるか・強さ λ", "7. L2・λ"],
        ["highOrder.termSets", "[]（なし）", "保存した項の組み合わせ（1行 = \"名前|式|k の番号 …\"）", "7. 組み合わせ"]],
        widths=(3.3, 2.2, 4.95, 1.68), row_h=0.36)
    point_bar(s, 6.12, "λ はそろえた係数に対する値（1 Shot 7×11 点の標準 0.1。点を減らすときは λ × √(77 / 点の数)）。"
              "highOrder は評価の結果を変えない（「続きから再開」の比較にも使わない）",
              icon="i", icon_color=BODY, h=0.7, size=12)
    notes(s, "「7. Shot高次補正」のタブの初期値。excludedTerms は「VIF で選ぶ」・表の「使う」・点の数による項の絞り込みで変わり、"
             "「設定を保存…」の JSON にも入る。termSets は画面の「保存」で増え、MATLAB の設定（getpref）にも残る。"
             "unit・λ・回帰・保存した組み合わせは「補正の比較」でまとめて並べて比べられる。outerMaker を none 以外にすると、"
             "外周の Shot だけその式で補正する（保存した組み合わせを選ぶと、その使わない項 outerExcludedTerms と名前 outerName も入る）。"
             "既定の設定ファイルにも書ける。"
             "計算の方法と式は「8. ヘルプ」の「Shot 高次補正の評価」。")

    # 15. options と画面の定数
    s = deck.body("03　設定の項目と初期値", "手法の細かい設定（methods.options）と画面の定数")
    rows = [["methods.options", "初期値"],
            ["gprMaxIterations", "300"], ["backfitMaxIterations ・ backfitToleranceNm", "20 ・ 0.001"],
            ["krigingLags", "6"], ["tpsLog10Lambda ・ rbfLog10Lambda", "[−8 3 45]（最小・最大・点数）"],
            ["rbfLengthScales", "[0.05 0.1 0.2 0.4 0.8 1.6]"], ["gamBasisAxis ・ gamBasisRadius", "8 ・ 8"],
            ["gamGpLengthScales", "[0.15 0.3 0.6]"], ["gamGpKnots ・ gamGpGridKnots", "10 ・ 7"]]
    table(s, 0.6, 1.5, [3.75, 2.3], rows, size=11.5, header_size=12, row_h=0.4)
    rows = [["画面などの定数（ファイル）", "値"],
            ["FIGURE_SIZE（画面の大きさ）", "[1440 900]"], ["FONT_SIZE ・ AXES_FONT_SIZE", "16 ・ 13"],
            ["DEFAULT_MARKER_SIZE（色の点）", "30"], ["MAX_CURVES（グラフの線の数）", "8"],
            ["MAX_MEASURE_PANELS ・ MAX_COMPARE_PANELS", "6 ・ 6"],
            ["SECONDS_PER_RUN ・ MATERN_NU_FACTOR", "2 秒 ・ 3 倍"],
            ["SECONDS_PER_SHOT_START（交換法）", "0.01 秒"], ["WORST_PERCENT（95%値）", "95"]]
    table(s, 6.85, 1.5, [3.98, 1.9], rows, size=11.5, header_size=12, row_h=0.43)
    point_bar(s, 5.65, "options は Python 版と同じ値。変えると推定の結果が変わり、Python 版との照合テスト"
              "（TestPythonReference）が合わなくなる", icon="!", icon_color=CORAL, h=0.75, size=13)
    notes(s, "options は精度と時間の兼ね合いで決めた値。変えるときは模擬データで比べてから。画面の定数は +appui/Core.m の先頭（画面はタブごとのファイルに分けてあり、定数は共通の Core.m にまとめてある）、"
             "時間の目安は estimateRunSeconds.m・exchangeSecondsGuide.m にある。")

    # 16. 編集の例
    s = deck.body("04　編集の例・注意・確かめ方", "編集の例 ― 同じ変更を設定ファイルとコードで書く")
    text(s, 0.6, 1.45, 12.13, 0.7, "例：D 最適も比べる・学習 Shot 数を 10〜50（10 刻み）・同心円を 60・100 mm・外周を 130 mm・"
         "異常値を MAD で除く・Fringe Zernike（次数 4）も比べる", size=13, color=INK, bold=True)
    code_box(s, 0.6, 2.25, 5.75, 4.45, "手順 A・B：設定ファイル（JSON）", [
        "{",
        '  "sampling": {',
        '    "method": ["ioptimal", "doptimal"],',
        '    "nShotsStart": 10, "nShotsEnd": 50,',
        '    "nShotsStep": 10,',
        '    "constraints": {',
        '      "radialZoneEdgesMm": [60, 100] }',
        "  },",
        '  "evaluation": { "outerRadiusMm": 130,',
        '    "excludeOutliers": true,',
        '    "outlierMethod": "mad" },',
        '  "methods": { "zernike": {',
        '    "enabled": true, "orders": 4 } }',
        "}"], color=CORAL, size=12.5)
    code_box(s, 6.6, 2.25, 6.13, 4.45, "手順 C：defaultSettings.m・recommendedMethods.m", [
        "% defaultSettings.m",
        'settings.sampling.method = ["ioptimal", "doptimal"];',
        "settings.sampling.nShotsStart = 10;",
        "settings.sampling.nShotsEnd = 50;",
        "settings.sampling.nShotsStep = 10;",
        "settings.sampling.nShots = alignest.shotCountList(10, 50, 10);",
        "settings.sampling.constraints.radialZoneEdgesMm = [60, 100];",
        "settings.evaluation.outerRadiusMm = 130;",
        "settings.evaluation.excludeOutliers = true;",
        'settings.evaluation.outlierMethod = "mad";',
        "% recommendedMethods.m（\"default\" の部分＝初期設定）",
        "methods.zernike.enabled = true;",
        "methods.zernike.orders = 4;"], color=PURPLE, size=11)
    notes(s, "JSON は書いた項目だけが変わる。コードでは学習 Shot 数の並び nShots も一緒に作る（shotCountList）。")

    # 17. 気をつけること
    s = deck.body("04　編集の例・注意・確かめ方", "気をつけること")
    tips = [("\"", TEAL, "型と形をそろえる", "文字は \"…\"（JSON も \"…\"）、並びは横の [a, b]。true・false は小文字。"
             "JSON は UTF-8 で保存する"),
            ("mm", CORAL, "単位を間違えない", "半径は mm、計測値・誤差は nm、計測率は %。半径 150 mm の Wafer が前提"),
            ("名", PURPLE, "項目の名前を変えない", "名前を変えると以前の設定ファイル・評価結果の Excel が読めなくなる。"
             "項目を足すときは defaultSettings.m に初期値を書く"),
            ("≠", GOLD, "結果が変わることを伝える", "推奨の組み合わせ・options を変えると評価の値が変わる。"
             "過去と比べるときは同じ設定ファイルで評価する"),
            ("✓", NAVY, "テストの期待値", "初期値を確かめるテスト（時間の目安・学習 Shot 数の並びなど）が失敗したら、"
             "意図した変更なら期待値を直す"),
            ("!", TEAL, "置き場所を確かめる", "既定の設定ファイルは、exe ならショートカットの作業フォルダーを探す。"
             "置いたら記録の「読み込みました」で確かめる")]
    for index, (glyph, color, heading, body) in enumerate(tips):
        row, col = divmod(index, 3)
        x = 0.6 + col * 4.145
        y = 1.55 + row * 2.6
        box(s, x, y, 3.84, 2.35, fill=WHITE, line=LINE)
        circle(s, x + 0.25, y + 0.25, 0.6, color, glyph, size=15)
        text(s, x + 1.0, y + 0.25, 2.7, 0.6, heading, size=15.5, color=INK, bold=True, anchor="m")
        text(s, x + 0.25, y + 1.0, 3.34, 1.3, body, size=12.5, color=BODY)
    notes(s, "設定ファイルは loadSettings が初期値の型に合わせて読む（文字の並び・数の行ベクトル）。知らない項目は無視する。")

    # 18. 確かめ方
    s = deck.body("04　編集の例・注意・確かめ方", "確かめ方 ― テスト・画面・Excel")
    steps(s, 0.6, 1.5, 6.6, [
        ("テストを全件流す（手順 C）", f"cd tests → runAllTests。「成功 {TEST_COUNT} 件、失敗 0 件、未完了 0 件」（2026-10-31 時点の件数）"),
        ("画面で初期値を見る", "「初期設定に戻す」を押し、「2. 評価の設定」の欄が新しい値か確かめる"),
        ("設定ファイルを開いて見る", "「設定を保存…」で保存した JSON の値（名前・単位・並び）"),
        ("模擬データで評価する", "sample_data の模擬データで評価し、結果の Excel の「設定」シートを見る"),
        ("exe でも同じことを確かめる", "作り直した exe（手順 C）、または既定の設定ファイルを置いたフォルダ（手順 B）で")],
        color=TEAL, gap=0.98)
    rows = [["変えたもの", "確かめるテスト"],
            ["設定の項目・読み込み", "TestSpatialModels・TestFilterExtendTiming"],
            ["学習 Shot の選び方・制約", "TestMeasurementPlan・TestLayoutDesignZernike"],
            ["比べる手法・options", "TestModelOptions・TestPythonReference"],
            ["Overlay 補正", "TestOverlay"],
            ["画面の定数・部品", "TestAppWorkflow"],
            ["既定の設定ファイル", "TestLayoutDesignZernike"]]
    table(s, 7.5, 1.5, [2.35, 2.88], rows, size=11.5, header_size=12, row_h=0.62)
    point_bar(s, 6.45, "すべて確かめてから配る。配った後に問題が出たら、前の設定ファイル・前の版に戻す",
              icon="✓", icon_color=TEAL, h=0.5, size=13)
    notes(s, "テストは matlab_app/tests にある。1つのファイルだけ確かめるときは runtests('TestOverlay') のように指定する。")

    # 19. まとめ
    s = deck.slide(dark=True)
    text(s, 0.8, 0.55, 12.13, 0.8, "まとめ", size=32, color=WHITE, bold=True)
    points = [(TEAL, "A 設定ファイルを配る：評価ごとに条件を変える。「設定を保存…」→「設定を読み込む…」"),
              (CORAL, "B 既定の設定ファイル（" + DEFAULTS_FILE + "）：exe を作り直さずに全員の初期値をそろえる"),
              (PURPLE, "C コードの初期値（defaultSettings.m・recommendedMethods.m）：テストと exe の作り直しが必要"),
              (YELLOW, "書いた項目だけが初期値に重なる。型・単位・項目の名前を保ち、変えたら画面と Excel で確かめる")]
    for index, (color, body) in enumerate(points):
        y = 1.7 + index * 1.22
        circle(s, 0.8, y + 0.12, 0.55, color, str(index + 1), label_color=INK if color == YELLOW else WHITE, size=16)
        text(s, 1.6, y, 6.2, 0.85, body, size=16, color=WHITE, anchor="m")
    box(s, 8.05, 1.6, 4.68, 5.0, fill=NAVY_PANEL)
    text(s, 8.35, 1.8, 4.1, 0.5, "関連する資料", size=17, color=WHITE, bold=True)
    related = [("README.md", "設定の項目・使い方の全体"),
               ("docs/admin_guide/admin_guide.html", "アプリの構成と保守の作業"),
               ("docs/windows_guide/windows_guide.html", "画面の使い方・exe を作る・配る"),
               ("「8. ヘルプ」タブ", "用語・カーネル・補正の式")]
    for index, (name, body) in enumerate(related):
        y = 2.5 + index * 0.98
        circle(s, 8.35, y + 0.08, 0.4, LIGHT, str(index + 1), label_color=INK, size=12)
        text(s, 8.95, y - 0.05, 3.6, 0.4, name, size=12.5, color=WHITE, bold=True)
        text(s, 8.95, y + 0.35, 3.6, 0.45, body, size=12, color=LIGHT2)
    text(s, 0.6, 7.08, 4.0, 0.25, DECK_NAME, size=9, color=LIGHT)
    text(s, 11.73, 7.08, 1.0, 0.25, str(deck.page), size=10, color=LIGHT, align="r")
    notes(s, "まず A・B を使い、全員の標準を恒久的に変えるときだけ C にする。どの方法でも、変えたら画面と Excel で確かめる。")

    properties = deck.prs.core_properties
    properties.title = DECK_TITLE
    properties.subject = "面内推定評価ツールの初期設定（評価の条件・比べる手法・画面の初期値）の変え方"
    properties.author = "alignment-estimation-benchmark"
    properties.last_modified_by = "alignment-estimation-benchmark"
    temporary = output.with_name(output.stem + ".building.pptx")
    deck.prs.save(str(temporary))
    os.replace(temporary, output)
    return deck.page


def main():
    output = Path(sys.argv[1]).resolve()
    reference = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_REFERENCE
    lock = output.with_name("~$" + output.name)
    if lock.exists():
        sys.exit(f"PowerPoint で開いています（{lock.name}）。閉じてからもう一度実行してください。")
    pages = build(output, reference)
    print(f"作成しました：{output}（{pages} ページ）")


if __name__ == "__main__":
    main()
