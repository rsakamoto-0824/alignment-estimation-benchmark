"""「加法カーネルと交互作用」の解説資料（PowerPoint）を作る。

使い方：python3 build_deck.py <出力 .pptx> [参照デッキ .pptx] [計算例の JSON]
- 参照デッキ：技術レポートPowerPoint（reference.pptx）。スライドマスター・テーマ・フォントをそのまま使い、
  元のスライドは消して、同じ色・位置・文字の大きさで新しいスライドを置く。
- 計算例の JSON：kernel_numbers.m（MATLAB）で、同梱の模擬データから計算した値。
出力は別名（.building.pptx）に書いてから置き換える。出力先に PowerPoint のロックファイル（~$…）があれば止める。
"""
from __future__ import annotations

import json
import math
import os
import sys
from pathlib import Path

from pptx import Presentation

from pptkit import (BODY, CORAL, CORAL_TINT, GOLD, GOLD_TINT, INK, LIGHT, LIGHT2, LINE, MUTED, NAVY, NAVY_PANEL,
                    PANEL, PURPLE, PURPLE_TINT, TEAL, TEAL_TINT, WHITE, YELLOW, background, bar_chart, box, chrome,
                    circle, doughnut, line, matrix, notes, point_bar, table, text)

HERE = Path(__file__).resolve().parent
DEFAULT_REFERENCE = Path.home() / ".claude/skills/artifact-template-powerpoint-2/assets/reference.pptx"
DECK_NAME = "加法カーネルと交互作用"
DECK_TITLE = "加法カーネルと交互作用の解説"   # ファイルの情報のタイトル
NOISE = "A9ACC6"   # ノイズの色（参照デッキの淡い紺）

# 項の色（全ページで同じ）：空間・Scan 方向・Mark 番号・Scan×Mark・ノイズ
TERM = {"space": TEAL, "scan": CORAL, "mark": PURPLE, "scan_mark": GOLD, "noise": NOISE}
TINT = {"space": TEAL_TINT, "scan": CORAL_TINT, "mark": PURPLE_TINT, "scan_mark": GOLD_TINT}


SMALL_NM = 0.01   # これより小さい τ は「≈0」と書く（下限付近まで縮んだ値）


def fmt(value, digits=2):
    return f"{value:.{digits}f}"


def signed(value):
    return "0.00" if abs(value) < 0.005 else f"{value:+.2f}"


def tau_text(value):
    return "≈0" if value < SMALL_NM else fmt(value)


class Deck:
    def __init__(self, reference: Path):
        self.prs = Presentation(str(reference))
        self.layout = self.prs.slides[0].slide_layout
        # 参照デッキのスライドを外す（マスター・テーマ・フォントは残る）
        id_list = self.prs.slides._sldIdLst
        for slide_id in list(id_list):
            self.prs.part.drop_rel(slide_id.rId)
            id_list.remove(slide_id)
        self.page = 0

    def slide(self, dark=False):
        slide = self.prs.slides.add_slide(self.layout)
        for placeholder in list(slide.placeholders):
            placeholder._element.getparent().remove(placeholder._element)
        background(slide, NAVY if dark else WHITE)
        self.page += 1
        return slide

    def body(self, chapter, title):
        slide = self.slide()
        chrome(slide, chapter, title, self.page, DECK_NAME)
        return slide


def formula_bar(slide, y, content, h=1.05, size=24):
    box(slide, 0.6, y, 12.13, h, fill=NAVY)
    text(slide, 0.6, y, 12.13, h, content, size=size, color=WHITE, font="Times New Roman", align="c", anchor="m")


def term_panel(slide, x, y, w, h, key, heading, formula, note, heading_size=17):
    box(slide, x, y, w, h, fill=TINT[key])
    text(slide, x + 0.25, y + 0.2, w - 0.5, 0.42, heading, size=heading_size, color=TERM[key], bold=True)
    text(slide, x + 0.15, y + 0.72, w - 0.3, 0.6, formula, size=19, color=INK, font="Times New Roman")
    text(slide, x + 0.25, y + 1.38, w - 0.5, h - 1.5, note, size=13, color=BODY)


def build(output: Path, reference: Path, numbers: dict):
    deck = Deck(reference)
    xi = numbers["X_additive_int"]
    xa = numbers["X_additive"]
    names = xi["tauNames"]
    tau = dict(zip(names, xi["tau"]))
    variance = {"space": xi["sigmaF"] ** 2, "scan": tau["scan"] ** 2, "mark": tau["mark"] ** 2,
                "scan_mark": tau["scan_mark"] ** 2, "noise": xi["sigmaN"] ** 2}
    total = sum(variance.values())

    # 1. 表紙
    s = deck.slide(dark=True)
    text(s, 0.8, 0.9, 6.5, 0.4, "技術解説資料 ｜ 面内推定 検証ツール（GPR）", size=14, color=LIGHT)
    text(s, 0.8, 1.9, 7.0, 1.0, "加法カーネルと交互作用", size=40, color=WHITE, bold=True)
    text(s, 0.8, 3.05, 6.6, 1.3, "GPR で質的変数（Scan 方向・Mark 番号）を\nどのようなカーネルで扱っているか", size=20,
         color=LIGHT2)
    text(s, 0.8, 6.35, 6.6, 0.35, "2026年10月 ｜ 計算例は同梱の模擬データを MATLAB R2022b で計算", size=12, color=LIGHT)
    text(s, 8.0, 0.75, 4.8, 0.35, "例：模擬データ Wafer 1・X 成分の分散の内訳", size=12, color=LIGHT)
    keys = ["space", "scan", "mark", "scan_mark", "noise"]
    labels = {"space": "空間 $σ_{f}^{2}$", "scan": "Scan 方向 $τ_{s}^{2}$", "mark": "Mark 番号 $τ_{m}^{2}$",
              "scan_mark": "Scan×Mark $τ_{sm}^{2}$", "noise": "ノイズ $σ_{n}^{2}$"}
    doughnut(s, 8.2, 1.15, 4.2, 3.6, [k for k in keys], [variance[k] / total for k in keys], [TERM[k] for k in keys])
    text(s, 9.3, 2.65, 2.0, 0.6, "$k(z,z)+σ_{n}^{2}$", size=18, color=WHITE, font="Times New Roman", align="c",
         anchor="m")
    for index, key in enumerate(keys):
        yy = 4.95 + index * 0.33
        circle(s, 8.55, yy + 0.07, 0.18, TERM[key])
        text(s, 8.85, yy, 3.9, 0.32, f"{labels[key]}　{variance[key] / total * 100:.0f}%　（{variance[key]:.2f} nm²）",
             size=13, color=WHITE)
    notes(s, "アプリの GPR で「加法カーネル」「加法カーネル＋交互作用」を選んだときに、カーネルがどんな式になっているかを説明する資料。"
             "右の図は、模擬データの1枚で推定した超パラメータから、1点の値の分散が各項にどれだけ割り振られたかを示す。")

    # 2. この資料で学ぶこと
    s = deck.body("", "この資料で学ぶこと")
    box(s, 0.6, 1.55, 12.13, 0.85, fill=PANEL)
    circle(s, 0.82, 1.69, 0.57, NAVY, "◎", size=16)
    text(s, 1.6, 1.55, 10.93, 0.85, "ゴール：アプリの「加法カーネル」「加法カーネル＋交互作用」が、どんな式のカーネルで、"
         "推定値にどう効くかを説明できる", size=16, color=INK, bold=True, anchor="m")
    cards = [("k", CORAL, "GPR と質的変数", "前提の整理。なぜ質的変数を入れるか、GPR の式、質的変数の8通りの扱い方", "p.3–5"),
             ("+", TEAL, "加法カーネル", "「同じ水準なら τ²」を足すカーネル。行列・変量効果・予測での効き方、固定効果との違い",
              "p.6–10"),
             ("×", PURPLE, "交互作用", "「両方が同じなら τ²」の項。全ペアの入れ方、平均に入れる交互作用・混合との違い",
              "p.11–14"),
             ("▶", GOLD, "計算例と使い方", "超パラメータの推定、模擬データでの数値例、アプリでの場所、使い分けと注意点",
              "p.15–21")]
    for index, (glyph, color, heading, body, pages) in enumerate(cards):
        x = 0.6 + index * 3.11
        box(s, x, 2.75, 2.81, 3.95, fill=WHITE, line=LINE)
        circle(s, x + 0.3, 3.05, 0.7, color, glyph, size=22, font="Times New Roman")
        text(s, x + 1.71, 3.05, 0.8, 0.7, f"0{index + 1}", size=30, color=LINE, bold=True, align="r", anchor="m")
        text(s, x + 0.3, 3.95, 2.21, 0.5, heading, size=19, color=INK, bold=True)
        text(s, x + 0.3, 4.55, 2.21, 1.5, body, size=13, color=BODY)
        text(s, x + 0.3, 6.15, 2.21, 0.3, pages, size=11, color=MUTED)
    notes(s, "4章構成。前提 → 加法カーネル → 交互作用 → 計算例と使い方の順に進む。")

    # 3. なぜ質的変数を入れるか
    s = deck.body("01　前提：GPR と質的変数", "なぜ質的変数を入れるか ― 計測値は3つの成分の足し算")
    parts = [("計測値 $y$", PANEL, INK), ("空間の滑らかな成分 $f(x)$", TEAL_TINT, TEAL),
             ("水準ごとのずれ", CORAL_TINT, CORAL), ("ノイズ $ε$", PANEL, MUTED)]
    xs = [0.6, 3.75, 6.9, 10.05]
    for index, (label, fill, color) in enumerate(parts):
        box(s, xs[index], 1.65, 2.68, 1.0, fill=fill)
        text(s, xs[index], 1.65, 2.68, 1.0, label, size=17, color=color if index else INK, bold=True, align="c", anchor="m")
        if index < 3:
            text(s, xs[index] + 2.68, 1.65, 0.47, 1.0, "=" if index == 0 else "+", size=26, color=MUTED, bold=True,
                 align="c", anchor="m", font="Times New Roman")
    text(s, 0.6, 2.8, 12.13, 0.5, "水準ごとのずれ：Wafer 上の位置では決まらず、「どの水準か」で決まる段差。空間の滑らかな関数だけでは表しにくい",
         size=14, color=BODY)
    factors = [("S", CORAL, "Scan 方向", "露光の走査の向き（1・2）。向きごとに全体が少しずれる"),
               ("M", PURPLE, "Mark 番号", "Shot 内の Mark の位置（1〜4 など）。位置ごとに決まったずれがある"),
               ("R", TEAL, "Shot 行の奇数・偶数", "Shot 配列の行の奇偶。選んだときだけ使う"),
               ("C", GOLD, "ファイルのその他の列", "Chuck など、標準の項目以外の列。選んだときだけ使う")]
    for index, (glyph, color, heading, body) in enumerate(factors):
        x = 0.6 + index * 3.11
        box(s, x, 3.45, 2.81, 2.05, fill=WHITE, line=LINE)
        circle(s, x + 0.25, 3.7, 0.6, color, glyph, size=18, font="Times New Roman")
        text(s, x + 1.0, 3.7, 1.7, 0.6, heading, size=16, color=INK, bold=True, anchor="m")
        text(s, x + 0.25, 4.45, 2.4, 1.0, body, size=13, color=BODY)
    point_bar(s, 5.75, "問い：水準ごとのずれを、平均の式に入れる（固定効果）か、カーネルに入れる（変量効果）か → 後者が「加法カーネル」",
              icon="?", icon_color=CORAL, fill=PANEL, size=15)
    notes(s, "アライメントの計測値には、Wafer 上の位置で滑らかに変わる成分のほかに、Scan 方向や Mark 番号で決まる段差がある。"
             "質的変数（カテゴリーの変数）として推定に入れると、段差を説明できる。")

    # 4. GPR の式
    s = deck.body("01　前提：GPR と質的変数", "GPR の式 ― 平均・カーネル・ノイズの3つで決まる")
    formula_bar(s, 1.6, "$y = h(x)^{T}β + f(x) + ε$ ,　 $f ~ GP(0, k(x, x′))$ ,　 $ε ~ N(0, σ_{n}^{2})$")
    panels = [("平均の部分", "$h(x)^{T}β$", "定数（と、固定効果にする質的変数の効果コード）。β は一般化最小二乗で求める", PANEL, BODY),
              ("カーネル", "$k(x, x′)$", "2点の値がどれだけ似ているか（共分散）。アプリの既定は matern52：$σ_{f}^{2}(1+√5r+5r^{2}/3)e^{−√5r}$、"
               "$r$ = 距離 / $ℓ$", TEAL_TINT, TEAL),
              ("ノイズ", "$σ_{n}^{2}$", "計測のばらつき。カーネルの対角に足す（$K_{y} = K + σ_{n}^{2}I$）", PANEL, BODY)]
    for index, (heading, formula, note, fill, color) in enumerate(panels):
        x = 0.6 + index * 4.145
        box(s, x, 2.95, 3.84, 2.5, fill=fill)
        text(s, x + 0.3, 3.2, 3.24, 0.45, heading, size=18, color=color, bold=True)
        text(s, x + 0.2, 3.75, 3.44, 0.6, formula, size=20, color=INK, font="Times New Roman")
        text(s, x + 0.3, 4.4, 3.24, 1.0, note, size=13, color=BODY)
    point_bar(s, 5.75, "予測：$ŷ_{*} = h_{*}^{T}β̂ + k_{*}^{T}K_{y}^{−1}(y − Hβ̂)$。超パラメータ（$ℓ$・$σ_{f}$・$τ$・$σ_{n}$）は"
              "対数周辺尤度を最大にして選ぶ", icon="k", icon_color=TEAL, size=15)
    notes(s, "GPR（ガウス過程回帰）は、平均の部分・カーネル・ノイズで決まる。質的変数は、平均の部分に入れる方法と、カーネルに入れる方法がある。")

    # 5. 8通りの扱い方
    s = deck.body("01　前提：GPR と質的変数", "質的変数の扱い方は8通り ― 平均に入れるか、カーネルに入れるか")
    rows = [["扱い方（id の末尾）", "平均の部分（固定効果）", "カーネル（変量効果）", "交互作用"],
            ["使わない（none）", "定数だけ", "空間だけ", "—"],
            ["同時推定（basis）", "定数 ＋ 効果コード", "空間だけ", "—"],
            ["Backfitting（backfit）", "質的変数の効果と GPR を交互に当てはめ直す", "空間だけ", "—"],
            ["同時推定＋交互作用（basis_int）", "＋ 効果コードの積（全ペア）", "空間だけ", "平均に入れる"],
            ["加法カーネル（additive）", "定数だけ", "空間 ＋ $τ^{2}$[同じ水準]（変数ごと）", "—"],
            ["加法カーネル＋交互作用（additive_int）", "定数だけ", "＋ $τ^{2}$[両方同じ]（全ペア）", "カーネルに入れる"],
            ["混合（mixed）", "定数 ＋ Mark 以外の効果コード", "空間 ＋ $τ_{m}^{2}$[Mark が同じ]", "—"],
            ["混合＋交互作用（mixed_int）", "定数 ＋ Mark 以外の効果コード", "＋ $τ^{2}$[両方同じ]（全ペア）", "カーネルに入れる"]]
    fills = [None, None, None, None, None, TEAL_TINT, TEAL_TINT, None, None]
    table(s, 0.6, 1.55, [3.65, 3.55, 3.33, 1.6], rows, fills=fills, size=12, row_h=0.47, bold_first=True)
    point_bar(s, 5.95, "この資料の主役は、色を付けた2行：加法カーネル（additive）と、加法カーネル＋交互作用（additive_int）",
              icon="!", icon_color=TEAL, fill=PANEL, h=0.8, size=15)
    notes(s, "アプリの「質的変数の扱い」の8通り。平均の部分に効果コードを入れるのが固定効果、カーネルに同じ水準の項を足すのが変量効果。"
             "[同じ水準] は、2点の水準が同じなら 1、違えば 0 という意味。")

    # 6. 加法カーネルの式
    s = deck.body("02　加法カーネル", "加法カーネルの式 ― 「同じ水準なら τ²」を足す")
    formula_bar(s, 1.55, "$k(z, z′) = σ_{f}^{2}ρ(r) + τ_{s}^{2}δ(s, s′) + τ_{m}^{2}δ(m, m′)$", size=27)
    text(s, 0.6, 2.7, 12.13, 0.4, "$z$ = (位置 $x$, Scan 方向 $s$, Mark 番号 $m$)　　$δ(a, b)$ = 1（$a = b$ のとき）、0（それ以外）",
         size=15, color=BODY, align="c")
    term_panel(s, 0.6, 3.2, 3.84, 2.35, "space", "空間の項", "$σ_{f}^{2}ρ(r)$",
               "位置が近いほど大きい。$ρ$ は matern52 などの形、$r$ = 距離 / $ℓ$")
    term_panel(s, 4.74, 3.2, 3.84, 2.35, "scan", "Scan 方向の項", "$τ_{s}^{2}δ(s, s′)$",
               "2点の Scan 方向が同じなら $τ_{s}^{2}$、違えば 0。距離には よらない")
    term_panel(s, 8.89, 3.2, 3.84, 2.35, "mark", "Mark 番号の項", "$τ_{m}^{2}δ(m, m′)$",
               "2点の Mark 番号が同じなら $τ_{m}^{2}$、違えば 0。Wafer の端と端でも同じ")
    point_bar(s, 5.85, "Shot 行の奇偶・その他の列を選ぶと、同じ形の項 $τ^{2}δ(…)$ が1つずつ加わる。ノイズは $K_{y} = K + σ_{n}^{2}I$",
              icon="+", icon_color=TEAL, size=15)
    notes(s, "加法カーネルは、空間のカーネルに、質的変数ごとの「水準が同じなら τ²」の項を足したもの。"
             "τ は変数ごとに1つの超パラメータで、空間の長さスケールなどと一緒に推定する。")

    # 7. 行列で見る
    s = deck.body("02　加法カーネル", "行列で見る ― 「同じ水準なら 1」の行列を重ねる")
    labels8 = ["A1", "A2", "A3", "A4", "B1", "B2", "B3", "B4"]
    scan8 = [1, 1, 1, 1, 2, 2, 2, 2]
    mark8 = [1, 2, 3, 4, 1, 2, 3, 4]
    marks_xy = [(-10, -12), (10, -12), (10, 12), (-10, 12)]
    points8 = [(mx, my) for mx, my in marks_xy] + [(mx + 26, my) for mx, my in marks_xy]
    def matern52(dmm, ell=45.0):
        r = dmm / ell
        return (1 + math.sqrt(5) * r + 5 / 3 * r * r) * math.exp(-math.sqrt(5) * r)
    space8 = [[matern52(math.dist(a, b)) for b in points8] for a in points8]
    same_scan = [[1.0 if scan8[i] == scan8[j] else 0.0 for j in range(8)] for i in range(8)]
    same_mark = [[1.0 if mark8[i] == mark8[j] else 0.0 for j in range(8)] for i in range(8)]
    text(s, 0.6, 1.5, 12.13, 0.4, "例：Shot A（Scan 1）と、隣の Shot B（Scan 2）の Mark 1〜4（A1〜B4 の8点）", size=14, color=BODY)
    cell = 0.29
    gx = [1.1, 4.55, 8.0]
    matrix(s, gx[0], 2.35, cell, space8, TEAL, labels8, title="$σ_{f}^{2}$ × 空間（近いほど濃い）", title_color=TEAL)
    matrix(s, gx[1], 2.35, cell, same_scan, CORAL, labels8, title="$τ_{s}^{2}$ × 同じ Scan 方向", title_color=CORAL)
    matrix(s, gx[2], 2.35, cell, same_mark, PURPLE, labels8, title="$τ_{m}^{2}$ × 同じ Mark 番号", title_color=PURPLE)
    for x in (gx[0] + 8 * cell + 0.08, gx[1] + 8 * cell + 0.08):
        text(s, x, 3.2, 0.5, 0.6, "+", size=28, color=MUTED, bold=True, align="c", anchor="m", font="Times New Roman")
    box(s, 10.75, 2.35, 1.98, 2.9, fill=PANEL)
    text(s, 10.85, 2.45, 1.8, 0.35, "読み方", size=13, color=INK, bold=True)
    text(s, 10.85, 2.85, 1.8, 2.35, ["対角（自分）はすべての項が効く", "A1 と B1：空間と Mark だけでつながる",
                                    "A1 と A2：空間と Scan だけでつながる"], size=12, color=BODY, space_after=4)
    point_bar(s, 5.85, "$K = σ_{f}^{2}K_{空間} + τ_{s}^{2}S + τ_{m}^{2}M$（$S$・$M$ は 0 と 1 の行列）。足しても正定値のまま（どの項も共分散になる形）",
              icon="∑", icon_color=TEAL, size=14)
    notes(s, "カーネル行列で見ると、空間の行列に、Scan 方向が同じ組だけ 1 の行列 S と、Mark 番号が同じ組だけ 1 の行列 M を、"
             "それぞれ τ² 倍して足したもの。S・M はグループごとに全部 1 のブロックなので正定値（半正定値）で、和もカーネルになる。"
             "空間の濃さは説明用に長さスケール 45 mm（matern52）で描いた。")

    # 8. 変量効果として読む
    s = deck.body("02　加法カーネル", "変量効果として読む ― 同じ水準の点は「同じずれ」を共有する")
    formula_bar(s, 1.55, "$y = β_{0} + f(x) + u_{s} + u_{m} + ε$ ,　 $u_{s} ~ N(0, τ_{s}^{2})$ ,　 $u_{m} ~ N(0, τ_{m}^{2})$", size=24)
    text(s, 0.6, 2.75, 6.1, 0.5, "共分散を計算すると、加法カーネルそのものになる", size=17, color=INK, bold=True)
    box(s, 0.6, 3.3, 6.1, 1.25, fill=PANEL)
    text(s, 0.7, 3.3, 5.9, 1.25, "$Cov(y_{i}, y_{j}) = σ_{f}^{2}ρ(r_{ij}) + τ_{s}^{2}δ(s_{i}, s_{j}) + τ_{m}^{2}δ(m_{i}, m_{j}) + σ_{n}^{2}δ_{ij}$",
         size=17, color=INK, font="Times New Roman", anchor="m", align="c")
    text(s, 0.6, 4.7, 6.1, 1.1, ["$u_{s}$：Scan 方向ごとに1つの値（Scan 1 のずれ、Scan 2 のずれ）",
                                "$u_{m}$：Mark 番号ごとに1つの値。同じ水準の点はすべて同じ値を足す",
                                "$f$・$u$・$ε$ は互いに独立と考える"], size=13, color=BODY, bullets=True, space_after=4)
    rows = [["", "τ が表すもの"], ["$τ_{s}$", "Scan 方向によるずれの大きさ（標準偏差、nm）"],
            ["$τ_{m}$", "Mark 番号によるずれの大きさ（標準偏差、nm）"], ["$τ → 0$", "その変数のずれは無い（項が消える）"]]
    table(s, 7.1, 2.75, [1.2, 4.43], rows, size=13, row_h=0.62, aligns=["c", "l"])
    point_bar(s, 5.85, "固定効果は「ずれの値」をそのまま推定する。変量効果は「ずれの大きさ τ」を推定し、ずれの値はデータから予測する",
              icon="u", icon_color=CORAL, fill=PANEL, size=15)
    notes(s, "加法カーネルは、水準ごとのずれを平均0・分散 τ² の確率変数（変量効果）として扱うことと同じ。"
             "同じ水準の点は同じずれを足すので、その分だけ共分散が τ² 大きくなる。")

    # 9. 予測でどう効くか
    s = deck.body("02　加法カーネル", "予測でどう効くか ― 3つの部分に分けられる")
    formula_bar(s, 1.55, "$ŷ_{*} = β̂_{0} + Σ_{i}[σ_{f}^{2}ρ(r_{*i}) + τ_{s}^{2}δ(s_{*}, s_{i}) + τ_{m}^{2}δ(m_{*}, m_{i})]α_{i}$ ,　"
                "$α = K_{y}^{−1}(y − β̂_{0})$", size=21)
    pieces = [("space", "空間の推定", "$Σ_{i} σ_{f}^{2}ρ(r_{*i}) α_{i}$", "近くの学習点の情報で決まる滑らかな成分"),
              ("scan", "Scan のずれの推定", "$û_{s} = τ_{s}^{2} Σ α_{i}$", "和は、Scan 方向が予測する点と同じ学習点について（同じ Scan の学習点すべてから求める）"),
              ("mark", "Mark のずれの推定", "$û_{m} = τ_{m}^{2} Σ α_{i}$", "和は、Mark 番号が予測する点と同じ学習点について（同じ Mark の学習点すべてから求める）")]
    for index, (key, heading, formula, note) in enumerate(pieces):
        term_panel(s, 0.6 + index * 4.145, 2.95, 3.84, 2.2, key, heading, formula, note)
    stats = [("縮む", "点が少ない水準・ノイズが大きいときは、ずれの推定が 0 の方へ縮む（$τ$ と $σ_{n}$ の比で決まる）"),
             ("0", "学習点に無い水準は $δ = 0$ なので、その項は効かない（平均に戻る）")]
    for index, (big, body) in enumerate(stats):
        x = 0.6 + index * 6.17
        box(s, x, 5.4, 5.96, 1.3, fill=WHITE, line=LINE)
        text(s, x + 0.2, 5.4, 1.4, 1.3, big, size=30, color=CORAL, bold=True, anchor="m", align="c")
        text(s, x + 1.7, 5.4, 4.1, 1.3, body, size=13, color=BODY, anchor="m")
    notes(s, "予測値は、空間の部分と、質的変数ごとの部分に分けられる。質的変数の部分は、同じ水準の学習点の α の和に τ² を掛けたもので、"
             "その水準のずれの予測（BLUP）になっている。固定効果と違い、情報が少ないと 0 に縮む。")

    # 10. 固定効果との違い
    s = deck.body("02　加法カーネル", "同時推定（固定効果）との違い")
    rows = [["観点", "同時推定（basis、固定効果）", "加法カーネル（additive、変量効果）"],
            ["効果の表し方", "効果コード（水準の数 − 1 列）を平均に入れる", "$τ^{2}$[同じ水準] をカーネルに足す"],
            ["推定するもの", "水準ごとの係数 β（一般化最小二乗）", "ずれの大きさ $τ$（尤度最大化）。ずれの値は予測する"],
            ["点が少ない水準", "そのまま推定する（ばらつきやすい）", "0 の方へ縮めて推定する（縮小推定）"],
            ["効果が無いとき", "係数は小さくなるが、推定の誤差の分だけ残る", "$τ$ が下限まで小さくなり、項がほぼ消える"],
            ["学習点に無い水準", "効果 0（全水準の平均として扱う）", "効果 0（その項が効かない）"],
            ["増えるもの", "平均の係数（水準の数に応じて）", "超パラメータ $τ$（変数ごとに1つ）"]]
    table(s, 0.6, 1.55, [2.6, 4.75, 4.78], rows, size=13, row_h=0.58, bold_first=True)
    point_bar(s, 5.85, "混合（mixed）は、Mark 番号だけ変量効果（$τ_{m}^{2}δ$）、Scan 方向などは固定効果（効果コード）にする中間の形",
              icon="≈", icon_color=PURPLE, fill=PANEL, size=15)
    notes(s, "同じ質的変数でも、平均に入れる（固定効果）か、カーネルに入れる（変量効果）かで、推定のふるまいが変わる。"
             "効果の有無がはっきりしない変数は、変量効果だと自動で弱められる。")

    # 11. 交互作用とは
    s = deck.body("03　交互作用", "交互作用とは ― 組み合わせでだけ現れるずれ")
    text(s, 0.6, 1.5, 12.13, 0.4, "例：Scan 方向（行）× Mark 番号（列）ごとのずれ [nm]（説明用に作った数値）", size=14, color=BODY)
    scan_effect = [1.0, -1.0]
    mark_effect = [1.5, -1.0, -1.5, 1.0]
    def grid_table(x, extra):
        rows = [["", "M1", "M2", "M3", "M4"]]
        fills = [None]
        for r in range(2):
            row = [f"Scan {r + 1}"]
            for c in range(4):
                value = scan_effect[r] + mark_effect[c] + extra.get((r, c), 0.0)
                row.append(f"{value:+.1f}")
            rows.append(row)
            fills.append(None)
        return rows, fills
    rows, fills = grid_table(0.6, {})
    table(s, 0.6, 2.05, [1.25, 1.05, 1.05, 1.05, 1.05], rows, size=15, row_h=0.6, aligns=["c"] * 5, bold_first=True)
    text(s, 0.6, 3.95, 5.45, 0.9, "主効果だけ：各セル = 行のずれ + 列のずれ\n（Scan を変えると、全 Mark が同じだけ動く）", size=13,
         color=BODY)
    rows, fills = grid_table(6.9, {(0, 2): 2.0})
    shape = table(s, 7.28, 2.05, [1.25, 1.05, 1.05, 1.05, 1.05], rows, size=15, row_h=0.6, aligns=["c"] * 5, bold_first=True)
    cell = shape.table.cell(1, 3)
    cell.fill.solid()
    from pptkit import rgb
    cell.fill.fore_color.rgb = rgb(GOLD_TINT)
    text(s, 7.28, 3.95, 5.45, 0.9, "交互作用あり：Scan 1 × M3 だけ +2.0 上乗せ\n（Scan を変えたときの動きが Mark で違う）", size=13,
         color=BODY)
    text(s, 6.2, 2.4, 0.9, 0.8, "→", size=30, color=MUTED, bold=True, align="c", anchor="m")
    point_bar(s, 5.1, "主効果（$u_{s}$・$u_{m}$）だけでは、表の各行を平行に動かすことしかできない。特定の組み合わせの上乗せには、"
              "組み合わせごとのずれ $u_{sm}$ が要る", icon="×", icon_color=GOLD, fill=GOLD_TINT, h=1.1, size=15)
    notes(s, "交互作用は、2つの質的変数の組み合わせでだけ現れるずれ。主効果の足し算では表せないので、組み合わせごとのずれを別に置く。")

    # 12. 交互作用の項の式
    s = deck.body("03　交互作用", "交互作用の項 ― 「両方が同じなら τ²」を足す")
    formula_bar(s, 1.55, "$k(z, z′) = σ_{f}^{2}ρ(r) + τ_{s}^{2}δ(s,s′) + τ_{m}^{2}δ(m,m′) + τ_{sm}^{2}δ(s,s′)δ(m,m′)$", size=24)
    both = [[same_scan[i][j] * same_mark[i][j] for j in range(8)] for i in range(8)]
    cell = 0.27
    matrix(s, 1.1, 3.25, cell, same_scan, CORAL, labels8, title="$S$：同じ Scan 方向", title_color=CORAL)
    text(s, 1.1 + 8 * cell + 0.05, 4.0, 0.5, 0.6, "⊙", size=24, color=MUTED, bold=True, align="c", anchor="m",
         font="Times New Roman")
    matrix(s, 4.05, 3.25, cell, same_mark, PURPLE, labels8, title="$M$：同じ Mark 番号", title_color=PURPLE)
    text(s, 4.05 + 8 * cell + 0.05, 4.0, 0.5, 0.6, "=", size=28, color=MUTED, bold=True, align="c", anchor="m",
         font="Times New Roman")
    matrix(s, 7.0, 3.25, cell, both, GOLD, labels8, title="$S⊙M$：両方が同じ（要素ごとの積）", title_color=GOLD)
    box(s, 9.65, 2.95, 3.08, 3.6, fill=GOLD_TINT)
    text(s, 9.8, 3.05, 2.8, 0.4, "変量効果で読むと", size=14, color=GOLD, bold=True)
    text(s, 9.8, 3.5, 2.8, 3.0, ["「同じ Scan で同じ Mark」の組ごとのずれ $u_{sm} ~ N(0, τ_{sm}^{2})$ を足す",
                                "この8点では Shot A と B の Scan が違うので、両方が同じなのは自分だけ",
                                "実際の Wafer では、同じ Scan 方向の別の Shot の同じ Mark どうし（例：A1 と、2つ隣の Shot の Mark 1）がつながる"],
         size=12, color=BODY, space_after=6)
    notes(s, "交互作用の項は、Scan 方向が同じ行列 S と Mark 番号が同じ行列 M の、要素ごとの積（アダマール積）に τ_sm² を掛けたもの。"
             "8点の例では Shot A と B の Scan が違うので対角だけが 1 になるが、実際の Wafer では同じ Scan 方向の別の Shot の同じ Mark どうしがつながる。")

    # 13. 全ペア
    s = deck.body("03　交互作用", "選んだ質的変数の全ペアを入れる（3つ以上の積は入れない）")
    formula_bar(s, 1.55, "$k = σ_{f}^{2}ρ(r) + Σ_{f} τ_{f}^{2}δ_{f} + Σ_{f<g} τ_{fg}^{2}δ_{f}δ_{g}$", size=26)
    rows = [["選んだ質的変数", "主効果の τ", "交互作用の τ（ペア）", "τ の合計"],
            ["Scan 方向・Mark 番号（初期設定）", "2", "1（Scan×Mark）", "3"],
            ["＋ Shot 行の奇偶", "3", "3（Scan×Mark、Scan×行、Mark×行）", "6"],
            ["＋ その他の列（例：Chuck）", "4", "6", "10"]]
    table(s, 0.6, 2.95, [4.2, 1.9, 4.1, 1.93], rows, size=13, row_h=0.55, aligns=["l", "c", "l", "c"], bold_first=True)
    cards = [("p", "変数が p 個なら", "主効果 p 個 ＋ 交互作用 p(p−1)/2 個の τ を推定する。多いほど最適化に時間がかかる"),
             ("1", "学習点で水準が1つの変数", "その変数と、その変数を含むペアは使わない（推定できないため）"),
             ("3", "3つ以上の積", "Scan×Mark×行 のような3変数の組み合わせは入れない")]
    for index, (glyph, heading, body) in enumerate(cards):
        x = 0.6 + index * 4.145
        box(s, x, 5.35, 3.84, 1.35, fill=WHITE, line=LINE)
        circle(s, x + 0.2, 5.55, 0.5, GOLD, glyph, size=15, font="Times New Roman")
        text(s, x + 0.85, 5.5, 2.9, 0.4, heading, size=14, color=INK, bold=True)
        text(s, x + 0.85, 5.9, 2.9, 0.8, body, size=12, color=BODY)
    notes(s, "アプリの「＋交互作用」は、選んだ質的変数（X・Y 成分ごとの選択）の全ペアについて、両方が同じなら τ² の項を足す。"
             "3変数以上の組み合わせは入れない。")

    # 14. 平均に入れる交互作用・混合
    s = deck.body("03　交互作用", "ほかの「＋交互作用」との違い")
    columns = [("同時推定＋交互作用", "basis_int", "平均に入れる（固定効果）",
                "$h = [1, c_{s}, c_{m}, c_{s}⊙c_{m}]$", "効果コードの列どうしの積を、平均の基底に加える。組み合わせごとの係数をそのまま推定する",
                PANEL, BODY),
               ("加法カーネル＋交互作用", "additive_int", "カーネルに入れる（変量効果）",
                "$+ τ_{sm}^{2}δ(s,s′)δ(m,m′)$", "主効果も交互作用もカーネルの項。平均は定数だけ", TEAL_TINT, TEAL),
               ("混合＋交互作用", "mixed_int", "主効果は混合、交互作用はカーネル",
                "$h = [1, c_{s}]$ ,　$+ τ_{m}^{2}δ_{m} + τ_{sm}^{2}δ_{s}δ_{m}$",
                "Mark 以外の主効果は効果コード（固定）、Mark の主効果と全ペアの交互作用はカーネル（変量）", PURPLE_TINT, PURPLE)]
    for index, (heading, ident, kind, formula, body, fill, color) in enumerate(columns):
        x = 0.6 + index * 4.145
        box(s, x, 1.6, 3.84, 4.0, fill=fill)
        text(s, x + 0.25, 1.75, 3.34, 0.45, heading, size=18, color=color, bold=True)
        text(s, x + 0.25, 2.2, 3.34, 0.35, ident, size=13, color=MUTED, font="Times New Roman")
        text(s, x + 0.25, 2.65, 3.34, 0.4, kind, size=14, color=INK, bold=True)
        text(s, x + 0.15, 3.15, 3.54, 0.9, formula, size=16, color=INK, font="Times New Roman")
        text(s, x + 0.25, 4.1, 3.34, 1.4, body, size=13, color=BODY)
    point_bar(s, 5.85, "$c_{s}$・$c_{m}$ は効果コード（和が 0 になる 0・±1 の列）、$⊙$ は列どうしの要素ごとの積。"
              "交互作用は平均に入れると係数が増え、カーネルに入れると $τ$ が1つ増える", icon="⊙", icon_color=GOLD, size=14)
    notes(s, "交互作用を平均の部分に入れる（basis_int）と、組み合わせごとの係数を固定効果として推定する。"
             "カーネルに入れる（additive_int・mixed_int）と、組み合わせごとのずれを変量効果として扱う。")

    # 15. 超パラメータの推定
    s = deck.body("04　計算例と使い方", "超パラメータの推定 ― τ も一緒に尤度で選ぶ")
    formula_bar(s, 1.55, "$−log L(θ) = ½(y − Hβ̂)^{T}K_{y}^{−1}(y − Hβ̂) + ½ log|K_{y}| + (n/2) log 2π$", size=23)
    rows = [["超パラメータ", "初期値", "探す範囲"],
            ["長さスケール $ℓ$", "説明変数の標準偏差", "$10^{−3}$ 〜 $10^{3}$"],
            ["空間の大きさ $σ_{f}$", "$std(y)/√2$", "$10^{−3}$ 〜 $10^{3}$ × $std(y)$"],
            ["質的変数の $τ$（項ごと）", "$0.2 × std(y)$", "$10^{−4}$ 〜 $10$ × $std(y)$"],
            ["ノイズ $σ_{n}$", "$std(y)/√2$", "$10^{−2}$ 〜 $10$ × $std(y)$"]]
    table(s, 0.6, 2.9, [3.3, 2.7, 3.0], rows, size=13, row_h=0.52)
    box(s, 9.9, 2.9, 2.83, 2.6, fill=PANEL)
    text(s, 10.05, 3.0, 2.6, 0.35, "探し方", size=13, color=INK, bold=True)
    text(s, 10.05, 3.4, 2.6, 2.05, ["対数の値 $θ$ = (log $ℓ$, log $σ_{f}$, log $τ_{1}$…, log $σ_{n}$) を、上下限つきの L-BFGS で動かす",
                                   "勾配：$∂K/∂ log τ_{k} = 2τ_{k}^{2}D_{k}$（$D_{k}$ は 0・1 の行列）"], size=12, color=BODY, space_after=5)
    point_bar(s, 5.85, "β（平均）は一般化最小二乗で消去して、超パラメータだけを探す。$τ$ が下限（$10^{−4}$ × $std(y)$）に"
              "張り付いたら「その変数のずれは無い」と読める", icon="θ", icon_color=TEAL, size=14)
    notes(s, "τ は、空間の長さスケール・大きさ・ノイズと一緒に、対数周辺尤度を最大にして決める。"
             "アプリでは Toolbox を使わずに、上下限つきの L-BFGS（alignest.optim.minimizeBox）で解析的な勾配を使って探す。")

    # 16. 計算例：推定した超パラメータ
    s = deck.body("04　計算例と使い方", "計算例 ① ― 模擬データで推定した超パラメータ")
    text(s, 0.6, 1.5, 12.13, 0.45, f"模擬データ（同梱の sample_alignment.csv）{numbers['unit']}・{numbers['plan']}（学習 {xi['nTrain']} 点）・"
         "GPR matern52・説明変数 X,Y,半径", size=14, color=BODY)
    def column(item, axis):
        t = dict(zip(item["tauNames"], item["tau"]))
        return [fmt(item["lengthScale"], 3), fmt(item["sigmaF"]), tau_text(t["scan"]), tau_text(t["mark"]),
                tau_text(t["scan_mark"]) if "scan_mark" in t else "—", fmt(item["sigmaN"]), fmt(item["rmse"])]
    ya = numbers["Y_additive"]
    yi = numbers["Y_additive_int"]
    heads = ["長さスケール $ℓ$（半径で割った単位）", "空間の大きさ $σ_{f}$ [nm]", "Scan 方向 $τ_{s}$ [nm]", "Mark 番号 $τ_{m}$ [nm]",
             "Scan×Mark $τ_{sm}$ [nm]", "ノイズ $σ_{n}$ [nm]", "評価点の RMSE [nm]"]
    cols = [column(xa, "X"), column(xi, "X"), column(ya, "Y"), column(yi, "Y")]
    rows = [["超パラメータ", "X：加法", "X：加法＋交互作用", "Y：加法", "Y：加法＋交互作用"]]
    for r, head in enumerate(heads):
        rows.append([head] + [c[r] for c in cols])
    fills = [None] * len(rows)
    fills[5] = GOLD_TINT
    table(s, 0.6, 2.05, [4.33, 1.95, 1.95, 1.95, 1.95], rows, size=13, row_h=0.43, aligns=["l", "c", "c", "c", "c"],
          fills=fills)
    yt = dict(zip(yi["tauNames"], yi["tau"]))
    text(s, 0.6, 5.55, 12.13, 0.3, f"≈0：{SMALL_NM} nm 未満（探す範囲の下限の近くまで小さくなった値）", size=11, color=MUTED)
    point_bar(s, 5.9, f"X：交互作用の無いデータなので $τ_{{sm}}$ ≈ 0。Y（加法＋交互作用）：$τ_{{s}}$・$τ_{{m}}$ が ≈0 になり、代わりに $τ_{{sm}}$ = "
              f"{yt['scan_mark']:.2f} nm。組み合わせごとのずれ（8 通り）は主効果も表せるので、どちらに割り振るかは尤度しだい",
              icon="1", icon_color=GOLD, fill=PANEL, h=0.95, size=14)
    notes(s, "同梱の模擬データは、Scan 方向の差と Mark 番号ごとのずれを入れて作り、交互作用は入れていない。"
             "X 成分では τ_sm が下限付近まで小さくなり、データに交互作用が無いことと合っている。"
             "Y 成分では τ_s・τ_m が小さくなり、τ_sm がずれを引き受けた。Scan 2 × Mark 4 の 8 通りの組み合わせごとのずれは、"
             "主効果の足し算も表せるので、どの項に割り振るかがデータだけでは決まりにくい（予測値はどちらでも近い）。"
             "値は Excel の「Waferごと」の超パラメータ（tau_scan_X・tau_mark_X・tau_scan_mark_X など）と同じ。")

    # 17. 計算例：2点の共分散の内訳
    s = deck.body("04　計算例と使い方", "計算例 ② ― 2点の共分散は、どの項でつながるか")
    example = xi["example"]
    text(s, 0.6, 1.5, 12.13, 0.45, f"基準の点：評価点の1つ（Scan {example['scan']:.0f}・Mark {example['mark']:.0f}、"
         f"x = {example['x']:.1f} mm, y = {example['y']:.1f} mm）と、近くの4種類の点（加法カーネル＋交互作用、X 成分）",
         size=14, color=BODY)
    rows = [["相手の点", "距離 [mm]", "空間 $σ_{f}^{2}ρ$", "Scan $τ_{s}^{2}δ$", "Mark $τ_{m}^{2}δ$", "Scan×Mark", "合計 [nm²]"]]
    for pair in xi["pairs"]:
        rows.append([pair["label"], fmt(pair["distanceMm"], 1), fmt(pair["spatial"]), fmt(pair["scan"]), fmt(pair["mark"]),
                     fmt(pair["scanMark"]), fmt(pair["total"])])
    rows.append(["（自分自身）", "0", fmt(xi["sigmaF"] ** 2), fmt(tau["scan"] ** 2), fmt(tau["mark"] ** 2),
                 fmt(tau["scan_mark"] ** 2), fmt(xi["self"])])
    table(s, 0.6, 2.05, [3.9, 1.25, 1.45, 1.45, 1.45, 1.3, 1.33], rows, size=13, row_h=0.5,
          aligns=["l", "c", "c", "c", "c", "c", "c"])
    point_bar(s, 5.75, "同じ Mark 番号の点は、離れていても $τ_{m}^{2}$ でつながる。Scan 方向が同じなら さらに $τ_{s}^{2}$。"
              "空間の項だけでは、離れた点どうしは ほぼ つながらない", icon="2", icon_color=GOLD, fill=PANEL, h=1.0, size=15)
    notes(s, "推定した超パラメータで、2点のカーネルの値（共分散）を項ごとに分けた。0 の欄は、その項の条件（同じ水準）を満たさないため。"
             "自分自身との値は、各項の分散の和（ノイズを除く）。")

    # 18. 計算例：予測の内訳
    s = deck.body("04　計算例と使い方", "計算例 ③ ― 1点の推定値の内訳と、推定した水準のずれ")
    parts = [("β̂₀（平均）", example["beta0"], NAVY), ("空間", example["spatial"], TEAL)]
    part_names = {"scan": ("Scan のずれ", CORAL), "mark": ("Mark のずれ", PURPLE), "scan_mark": ("Scan×Mark", GOLD)}
    for name, value in zip(names, example["parts"]):
        parts.append((part_names[name][0], value, part_names[name][1]))
    x = 0.6
    width = 1.55
    for index, (label, value, color) in enumerate(parts):
        box(s, x, 1.65, width, 1.35, fill=WHITE, line=LINE)
        text(s, x, 1.72, width, 0.45, label, size=12, color=MUTED, bold=True, align="c")
        text(s, x, 2.15, width, 0.7, signed(value), size=26, color=color, bold=True, align="c", anchor="m")
        x += width
        text(s, x, 1.65, 0.35, 1.35, "+" if index < len(parts) - 1 else "=", size=22, color=MUTED, bold=True, align="c",
             anchor="m")
        x += 0.35
    box(s, x, 1.65, 1.95, 1.35, fill=NAVY)
    text(s, x, 1.72, 1.95, 0.45, "推定値 [nm]", size=12, color=LIGHT, bold=True, align="c")
    text(s, x, 2.15, 1.95, 0.7, signed(example["prediction"]), size=26, color=WHITE, bold=True, align="c", anchor="m")
    text(s, 0.6, 3.1, 12.13, 0.4, f"（この点の実測値は {example['measured']:+.2f} nm。単位 nm、X 成分）", size=13, color=MUTED)
    levels = xi["levelEffects"]
    mark_random = levels["mark"]["effects"]
    scan_random = levels["scan"]["effects"]
    # 同時推定（固定効果）の係数：[定数, Scan の効果コード 1 列, Mark の効果コード 3 列]。最後の水準は他の和の符号違い
    beta = numbers["X_basis"]["beta"]
    scan_fixed = [beta[1], -beta[1]]
    mark_fixed = beta[2:5] + [-sum(beta[2:5])]
    rows = [["Mark 番号のずれ [nm]", "M1", "M2", "M3", "M4"],
            ["加法カーネル（変量効果の予測 $û_{m}$）"] + [signed(v) for v in mark_random],
            ["同時推定（固定効果の係数）"] + [signed(v) for v in mark_fixed]]
    table(s, 0.6, 3.65, [3.75, 1.07, 1.07, 1.07, 1.07], rows, size=13, row_h=0.5, aligns=["l", "c", "c", "c", "c"])
    rows = [["Scan 方向のずれ [nm]", "Scan 1", "Scan 2"],
            ["加法カーネル $û_{s}$"] + [signed(v) for v in scan_random],
            ["同時推定"] + [signed(v) for v in scan_fixed]]
    table(s, 9.15, 3.65, [1.8, 0.89, 0.89], rows, size=12, row_h=0.5, aligns=["l", "c", "c"])
    big = [k for k in range(4) if abs(mark_fixed[k]) > 0.3]
    mark_ratio = sum(mark_random[k] / mark_fixed[k] for k in big) / len(big)
    scan_ratio = scan_random[0] / scan_fixed[0]
    point_bar(s, 5.55, f"変量効果の予測は、固定効果の係数より 0 に近い（この例で Mark は約 {mark_ratio * 10:.0f} 割、"
              f"Scan は約 {scan_ratio * 10:.0f} 割の大きさ）。符号と並びは同じ。$τ$ が小さい変数ほど強く縮むのが目安",
              icon="3", icon_color=GOLD, fill=PANEL, h=1.15, size=14)
    notes(s, "推定値 = 平均 + 空間 + 質的変数の各項。質的変数の項は、その水準の学習点の α の和に τ² を掛けた値（変量効果の予測、BLUP）。"
             "同じ学習点で同時推定（固定効果）にしたときの効果の値と比べると、符号と並びは同じで、0 の方へ縮んでいる。"
             "τ が小さい Scan 方向ほど強く縮む。固定効果の値は別のモデル（超パラメータも別に推定）なので、比はおおよその目安。")

    # 19. 計算例：手法の比較
    s = deck.body("04　計算例と使い方", "計算例 ④ ― 質的変数の扱い方で、推定誤差はどう変わるか")
    compare = numbers["compare"]
    order = ["none", "basis", "basis_int", "additive", "additive_int", "mixed", "mixed_int"]
    short = {"none": "使わない", "basis": "同時推定", "basis_int": "同時推定\n＋交互作用", "additive": "加法カーネル",
             "additive_int": "加法カーネル\n＋交互作用", "mixed": "混合", "mixed_int": "混合\n＋交互作用"}
    modes = compare["methods"]
    rmse = compare["rmse"]
    series = []
    for g, group in enumerate(compare["groups"]):
        series.append((f"学習 {group}", [rmse[modes.index(m)][g] for m in order]))
    bar_chart(s, 0.6, 1.5, 8.4, 5.2, [short[m] for m in order], series, [NAVY, TEAL],
              value_axis_title="RMSE（X・Y 合計）[nm]、Wafer 5 枚の平均", font_size=12)
    best15 = min(order, key=lambda m: rmse[modes.index(m)][0])
    box(s, 9.3, 1.6, 3.43, 5.0, fill=PANEL)
    text(s, 9.45, 1.7, 3.15, 0.35, "読み方", size=13, color=INK, bold=True)
    text(s, 9.45, 2.1, 3.15, 4.4, ["質的変数を使わないと、Scan・Mark のずれが残り、誤差が大きい",
                                  "固定効果（同時推定）と変量効果（加法カーネル）は、どちらもずれを取り込める",
                                  "交互作用の無いデータなので、＋交互作用では良くならない（推定する $τ$ が増える分、少し悪くなることもある）",
                                  f"この例で 15 Shot の最小は「{short[best15].replace(chr(10), '')}」"],
         size=12, color=BODY, space_after=6)
    notes(s, "同梱の模擬データ（5枚）で、GPR matern52（X,Y,半径）の質的変数の扱い方を変えて、評価点の RMSE（X・Y 合計）の Wafer 平均を比べた。"
             "学習 Shot は制約付き I 最適で 15・30 Shot。模擬データには交互作用が無いので、交互作用の項の効果は小さい。"
             "実データで交互作用がある場合は結果が変わるので、アプリの結果一覧で確かめる。")

    # 20. アプリでの場所
    s = deck.body("04　計算例と使い方", "アプリでの場所 ― どの関数が、どの項を作るか")
    rows = [["役割", "関数（+alignest/+models）", "中身"],
            ["扱い方の一覧", "gprFactorModes", "8通りの名前と画面の表示名"],
            ["項（0・1 の行列）を作る", "gprModeSetup（factorDesigns・component）", "主効果は [水準が同じ]、交互作用は2つの行列の要素ごとの積"],
            ["カーネルを計算する", "gprKernel", "空間のカーネル ＋ Σ $τ^{2}$ × [同じ水準] の行列。勾配も"],
            ["初期値・範囲", "gprInitialParams", "$τ$ は 0.2 × std(y) から、$10^{−4}$〜$10$ × std(y) の範囲"],
            ["尤度と勾配", "gprObjective", "β を消去した負の対数周辺尤度"],
            ["予測", "gprPredict", "$Hβ̂ + K_{*}α$（$K_{*}$ に質的変数の項も入る）"]]
    table(s, 0.6, 1.55, [2.7, 4.15, 5.28], rows, size=12, row_h=0.5, bold_first=True)
    cards = [("画面", "「2. 評価の設定」の GPR の「質的変数の扱い」で、加法カーネル・加法カーネル＋交互作用を選ぶ（複数選ぶと比べる）"),
             ("質的変数", "使う変数は「質的変数」で選ぶ（X 成分・Y 成分で別にもできる）。Shot 行の奇偶・その他の列も同じ形の項になる"),
             ("Excel", "「Waferごと」の「超パラメータ」に tau_scan_X・tau_mark_X・tau_scan_mark_X（Y も）として出る")]
    for index, (heading, body) in enumerate(cards):
        x = 0.6 + index * 4.145
        box(s, x, 5.3, 3.84, 1.4, fill=WHITE, line=LINE)
        text(s, x + 0.2, 5.38, 3.44, 0.35, heading, size=14, color=INK, bold=True)
        text(s, x + 0.2, 5.75, 3.44, 0.9, body, size=11.5, color=BODY)
    notes(s, "加法カーネルの項は gprModeSetup で 0・1 の行列として作り、gprKernel で τ² を掛けて空間のカーネルに足す。"
             "推定した τ は Excel の超パラメータの欄で確かめられる。")

    # 21. 使い分けと注意点
    s = deck.body("04　計算例と使い方", "使い分けと注意点")
    tips = [("τ", TEAL, "効かない変数", "$τ$ が下限へ小さくなり、自動で弱まる。ただし $τ$ の数だけ計算が増える"),
            ("1", CORAL, "水準が1つ", "学習点で水準が1つしかない変数（Scan 方向が1種類など）は使わない"),
            ("=", PURPLE, "同じ分け方の変数", "Scan 方向と Shot 行の奇偶が同じ分け方だと、同じ行列が2つ。どちらに割り振るか決まらない（予測は同じ）"),
            ("×", GOLD, "主効果と交互作用の重なり", "組み合わせごとのずれは主効果も表せる。$τ$ の割り振りが変わっても、予測は近い（計算例①の Y）"),
            ("∅", NAVY, "学習に無い水準・組み合わせ", "その項は効かず、平均に戻る。交互作用は、組み合わせが学習 Shot にそろっていると効く"),
            ("✓", TEAL, "選び方", "結果一覧の RMSE・基準の手法との比較で、使わない・同時推定・加法カーネル（＋交互作用）を比べる")]
    for index, (glyph, color, heading, body) in enumerate(tips):
        row, col = divmod(index, 3)
        x = 0.6 + col * 4.145
        y = 1.6 + row * 2.6
        box(s, x, y, 3.84, 2.35, fill=WHITE, line=LINE)
        circle(s, x + 0.25, y + 0.25, 0.6, color, glyph, size=18, font="Times New Roman")
        text(s, x + 1.0, y + 0.25, 2.7, 0.6, heading, size=16, color=INK, bold=True, anchor="m")
        text(s, x + 0.25, y + 1.0, 3.34, 1.25, body, size=13, color=BODY)
    notes(s, "加法カーネルは効かない変数を自動で弱められるが、超パラメータが増えるので計算時間は増える。"
             "同じ分け方の変数を2つ選ぶと、どちらの τ に割り振るかが決まらない（予測値は変わらない）。")

    # 22. まとめ
    s = deck.slide(dark=True)
    text(s, 0.8, 0.55, 12.13, 0.8, "まとめと次のステップ", size=32, color=WHITE, bold=True)
    points = [(CORAL, "加法カーネル = 空間のカーネル ＋ 変数ごとの $τ^{2}δ$（同じ水準なら $τ^{2}$）"),
              (TEAL, "水準ごとのずれを変量効果として扱う。ずれは予測され、情報が少ないと 0 へ縮む"),
              (PURPLE, "＋交互作用 = 全ペアの $τ_{fg}^{2}δ_{f}δ_{g}$（両方が同じなら $τ^{2}$）。3変数の積は入れない"),
              (YELLOW, "$τ$ は空間の超パラメータと一緒に尤度で推定。主効果と交互作用でずれを分け合うことがあるので、$τ$ はまとめて読む")]
    for index, (color, body) in enumerate(points):
        y = 1.75 + index * 1.18
        circle(s, 0.8, y + 0.12, 0.55, color, str(index + 1), label_color=INK if color == YELLOW else WHITE, size=16)
        text(s, 1.6, y, 5.9, 0.8, body, size=17, color=WHITE, anchor="m")
    box(s, 7.75, 1.6, 4.98, 5.0, fill=NAVY_PANEL)
    circle(s, 8.05, 1.85, 0.65, TEAL, "▶", size=18)
    text(s, 8.9, 1.85, 3.58, 0.65, "次のステップ：実データで確かめる", size=17, color=WHITE, bold=True, anchor="m")
    steps = ["結果一覧で、使わない・同時推定・加法カーネル（＋交互作用）の RMSE を比べる",
             "Excel の tau_scan・tau_mark・tau_scan_mark を見て、効いている変数を確かめる",
             "$τ_{sm}$ が大きく RMSE も良くなれば、交互作用を残す",
             "効かない変数は外して、計算を軽くする"]
    for index, body in enumerate(steps):
        y = 2.85 + index * 0.92
        circle(s, 8.1, y + 0.1, 0.42, LIGHT, str(index + 1), label_color=INK, size=13)
        text(s, 8.75, y - 0.04, 3.73, 0.75, body, size=13.5, color=LIGHT2, anchor="m")
    text(s, 0.6, 7.08, 4.0, 0.25, DECK_NAME, size=9, color=LIGHT)
    text(s, 11.73, 7.08, 1.0, 0.25, str(deck.page), size=10, color=LIGHT, align="r")
    notes(s, "加法カーネルと交互作用の要点。実データでは、交互作用の有無はデータ次第なので、アプリの結果一覧と超パラメータで確かめる。")

    # 23. 参考文献
    s = deck.body("", "参考文献")
    references = [
        "[1]  C. E. Rasmussen, C. K. I. Williams (2006). Gaussian Processes for Machine Learning. MIT Press.（4章：カーネルの和・積）",
        "[2]  P. Z. G. Qian, H. Wu, C. F. J. Wu (2008). Gaussian process models for computer experiments with qualitative and "
        "quantitative factors. Technometrics, 50(3), 383–396.",
        "[3]  O. Roustant, E. Padonou, Y. Deville, A. Clément, G. Perrin, J. Giorla, H. Wynn (2020). Group kernels for Gaussian "
        "process metamodels with categorical inputs. SIAM/ASA Journal on Uncertainty Quantification, 8(2), 775–806.",
        "[4]  G. K. Robinson (1991). That BLUP is a good thing: the estimation of random effects. Statistical Science, 6(1), 15–32.",
        "[5]  D. Duvenaud (2014). Automatic Model Construction with Gaussian Processes. PhD thesis, University of Cambridge.",
        "[6]  MathWorks. fitrgp（Gaussian process regression model）documentation. MATLAB R2022b.",
    ]
    text(s, 0.6, 1.5, 12.13, 4.4, references, size=11.5, color=INK, font="Arial", space_after=8)
    box(s, 0.6, 6.05, 12.13, 0.7, fill=PANEL)
    circle(s, 0.8, 6.13, 0.54, BODY, "i", size=15, font="Times New Roman")
    text(s, 1.55, 6.05, 11.03, 0.7, "本資料の式は、アプリの実装（matlab_app/+alignest/+models の gprModeSetup・gprKernel・gprInitialParams・"
         "gprObjective）に合わせた。計算例は同梱の模擬データを MATLAB R2022b で計算した値。", size=12, color=INK, anchor="m")
    notes(s, "カーネルの和で質的変数を扱う考え方は [1][2][3]、変量効果の予測（BLUP）は [4]、加法カーネルは [5] を参照。")

    # ファイルの情報（タイトルなど）は参照デッキのものが残るので、この資料のものにする
    properties = deck.prs.core_properties
    properties.title = DECK_TITLE
    properties.subject = "アプリの GPR の質的変数の扱い（加法カーネル・交互作用）の解説"
    properties.author = "alignment-estimation-benchmark"
    properties.last_modified_by = "alignment-estimation-benchmark"
    temporary = output.with_name(output.stem + ".building.pptx")
    deck.prs.save(str(temporary))
    os.replace(temporary, output)
    return deck.page


def main():
    output = Path(sys.argv[1]).resolve()
    reference = Path(sys.argv[2]) if len(sys.argv) > 2 else DEFAULT_REFERENCE
    numbers_path = Path(sys.argv[3]) if len(sys.argv) > 3 else HERE / "kernel_numbers.json"
    lock = output.with_name("~$" + output.name)
    if lock.exists():
        sys.exit(f"PowerPoint で開いています（{lock.name}）。閉じてからもう一度実行してください。")
    numbers = json.loads(numbers_path.read_text(encoding="utf-8"))
    pages = build(output, reference, numbers)
    print(f"作成しました：{output}（{pages} ページ）")


if __name__ == "__main__":
    main()
