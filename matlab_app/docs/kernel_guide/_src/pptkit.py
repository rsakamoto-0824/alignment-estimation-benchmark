"""技術レポートPowerPoint（reference.pptx）の見た目で、図形・文字・表・行列を置く小さな部品集。

座標はすべてインチ（スライドは 13.333 × 7.5）。色・文字の大きさ・位置は参照デッキの値に合わせる。
数式は「$...$」で囲み、その中の _{…} を下付き、^{…} を上付きにする（Times New Roman）。
"""
from __future__ import annotations

import copy
import re

from lxml import etree
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# 参照デッキの色
NAVY = "22223B"
NAVY_PANEL = "34365A"
INK = "22223B"
BODY = "4A4E69"
MUTED = "6C6F7F"
LINE = "DADDE6"
PANEL = "F3F4F8"
WHITE = "FFFFFF"
LIGHT = "A9ACC6"
LIGHT2 = "D9DBEA"
CORAL, CORAL_TINT = "E76F51", "FBE3DC"
TEAL, TEAL_TINT = "2A9D8F", "D8F0ED"
PURPLE, PURPLE_TINT = "8D6A9F", "ECE3F1"
GOLD, GOLD_TINT = "9C7412", "FBF1D6"
YELLOW = "E9C46A"
JP = "BIZ UDGothic"
MATH = "Times New Roman"
CORNER_IN = 0.08          # 角丸の半径（参照デッキの角丸と同じ大きさ）
SUB_BASELINE = "-40000"   # 下付きの位置（参照デッキと同じ）
SUP_BASELINE = "30000"

_TOKEN = re.compile(r"(_\{[^}]*\}|\^\{[^}]*\}|_.|\^.)")


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def _runs(text: str):
    """「$数式$」と下付き・上付きを、(文字, 数式か, 位置) の並びにする"""
    parts = text.split("$")
    runs = []
    for index, part in enumerate(parts):
        if part == "":
            continue
        math = index % 2 == 1
        if not math:
            runs.append((part, False, None))
            continue
        for token in _TOKEN.split(part):
            if token == "":
                continue
            if token.startswith("_"):
                runs.append((token[2:-1] if token.startswith("_{") else token[1:], True, "sub"))
            elif token.startswith("^"):
                runs.append((token[2:-1] if token.startswith("^{") else token[1:], True, "sup"))
            else:
                runs.append((token, True, None))
    return runs


def _set_font(run, size, color, bold, math, font):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    name = MATH if math else font
    run.font.name = name
    rpr = run.font._rPr
    for tag in ("a:ea", "a:cs"):
        element = rpr.find(qn(tag))
        if element is None:
            element = etree.SubElement(rpr, qn(tag))
        element.set("typeface", name)


def fill_paragraph(paragraph, text, size, color, bold=False, font=JP, align="l"):
    paragraph.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
    for content, math, position in _runs(text):
        run = paragraph.add_run()
        run.text = content
        _set_font(run, size, color, bold, math, font)
        if position == "sub":
            run.font._rPr.set("baseline", SUB_BASELINE)
        elif position == "sup":
            run.font._rPr.set("baseline", SUP_BASELINE)


def text(slide, x, y, w, h, content, size=13, color=BODY, bold=False, font=JP, align="l", anchor="t",
         bullets=False, space_after=0, line_spacing=None, margin=0.05):
    """文字の箱。content は文字列（改行で段落）または段落の並び"""
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    paragraphs = content if isinstance(content, list) else content.split("\n")
    for index, line in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        fill_paragraph(paragraph, line, size, color, bold, font, align)
        if space_after:
            paragraph.space_after = Pt(space_after)
        if line_spacing:
            paragraph.line_spacing = line_spacing
        if bullets:
            _bullet(paragraph, size)
    return shape


def _bullet(paragraph, size):
    ppr = paragraph._p.get_or_add_pPr()
    indent = int(Pt(size).emu * 1.1)
    ppr.set("marL", str(indent))
    ppr.set("indent", str(-indent))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for element in ppr.findall(qn(tag)):
            ppr.remove(element)
    char = etree.SubElement(ppr, qn("a:buChar"))
    char.set("char", "•")


def box(slide, x, y, w, h, fill=None, line=None, rounded=True, line_width=0.75):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        shape.adjustments[0] = min(0.5, CORNER_IN / max(min(w, h), 1e-6))
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    shape.text_frame.text = ""
    return shape


def circle(slide, x, y, d, fill, label="", label_color=WHITE, size=14, font=JP, bold=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.fill.background()
    shape.shadow.inherit = False
    if label:
        frame = shape.text_frame
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        fill_paragraph(frame.paragraphs[0], label, size, label_color, bold, font, "c")
    return shape


def line(slide, x1, y1, x2, y2, color=MUTED, width=1.25, arrow=True):
    connector = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    connector.line.color.rgb = rgb(color)
    connector.line.width = Pt(width)
    if arrow:
        ln = connector.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")
    return connector


def chrome(slide, chapter, title, page, deck_name, dark=False):
    """章の見出し・題名・下の文書名とページ番号（参照デッキと同じ位置）"""
    if chapter:
        text(slide, 0.6, 0.30, 8.0, 0.30, chapter, size=12, color=MUTED, bold=True)
    text(slide, 0.6, 0.62, 12.13, 0.75, title, size=28, color=WHITE if dark else INK, bold=True, anchor="m")
    footer_color = LIGHT if dark else MUTED
    text(slide, 0.6, 7.08, 4.0, 0.25, deck_name, size=9, color=footer_color)
    text(slide, 11.73, 7.08, 1.0, 0.25, str(page), size=10, color=footer_color, align="r")


def background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def point_bar(slide, y, content, icon="✓", icon_color=TEAL, fill=WHITE, line_color=LINE, h=0.95, size=15):
    """下のまとめの帯（丸の記号＋太字の一文）"""
    box(slide, 0.6, y, 12.13, h, fill=fill, line=line_color if fill == WHITE else None)
    circle(slide, 0.82, y + (h - 0.57) / 2, 0.57, icon_color, icon, size=16)
    text(slide, 1.6, y, 10.93, h, content, size=size, color=INK, bold=True, anchor="m")


def table(slide, x, y, widths, rows, header_fill=BODY, size=12, header_size=13, row_h=0.42, fills=None,
          aligns=None, bold_first=False, math_font=False):
    """参照デッキの表（濃い灰色の見出し・細い枠線・白地）。fills[r] で行の背景を変えられる"""
    n_rows, n_cols = len(rows), len(widths)
    shape = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(sum(widths)), Inches(row_h * n_rows))
    tbl = shape.table
    tbl_pr = tbl._tbl.tblPr
    for attribute in ("firstRow", "bandRow"):
        tbl_pr.set(attribute, "0")
    style = tbl_pr.find(qn("a:tableStyleId"))
    if style is not None:
        tbl_pr.remove(style)
    for c, width in enumerate(widths):
        tbl.columns[c].width = Inches(width)
    for r in range(n_rows):
        tbl.rows[r].height = Inches(row_h)
        for c in range(n_cols):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Emu(109728)
            cell.margin_top = cell.margin_bottom = Emu(45720)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            header = r == 0
            fill = header_fill if header else (fills[r] if fills and fills[r] else WHITE)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(fill)
            align = "c" if header else (aligns[c] if aligns else "l")
            frame = cell.text_frame
            frame.word_wrap = True
            for index, line_text in enumerate(str(rows[r][c]).split("\n")):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                fill_paragraph(paragraph, line_text, header_size if header else size, WHITE if header else INK,
                               header or (bold_first and c == 0), JP, align)
            _cell_border(cell)
    return shape


def _cell_border(cell, color=LINE, width=9525):
    tc_pr = cell._tc.get_or_add_tcPr()
    tags = ("a:lnL", "a:lnR", "a:lnT", "a:lnB")
    for tag in tags:
        for element in tc_pr.findall(qn(tag)):
            tc_pr.remove(element)
    # 枠線は塗りより前に、左・右・上・下の順に置く（スキーマの順番）
    for index, tag in enumerate(tags):
        ln = etree.Element(qn(tag), w=str(width), cap="flat", cmpd="sng", algn="ctr")
        solid = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(solid, qn("a:srgbClr"), val=color)
        etree.SubElement(ln, qn("a:prstDash"), val="solid")
        tc_pr.insert(index, ln)


def matrix(slide, x, y, cell, values, colors, labels, title=None, title_color=INK, label_size=8, zero=PANEL):
    """n × n の行列を色の升目で描く。values[i][j] は 0〜1（0 は zero の色、1 は colors の色、間は混ぜる）"""
    n = len(values)
    for i in range(n):
        text(slide, x - 0.42, y + i * cell, 0.40, cell, labels[i], size=label_size, color=MUTED, align="r", anchor="m",
             margin=0.0)
        text(slide, x + i * cell, y - 0.26, cell, 0.24, labels[i], size=label_size, color=MUTED, align="c", anchor="b",
             margin=0.0)
        for j in range(n):
            value = values[i][j]
            color = _mix(zero, colors, value)
            square = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + j * cell), Inches(y + i * cell),
                                            Inches(cell), Inches(cell))
            square.fill.solid()
            square.fill.fore_color.rgb = rgb(color)
            square.line.color.rgb = rgb(WHITE)
            square.line.width = Pt(0.75)
            square.shadow.inherit = False
    if title:
        text(slide, x - 0.45, y + n * cell + 0.08, n * cell + 0.45, 0.6, title, size=12, color=title_color, bold=True,
             align="c")


def _mix(low, high, t):
    t = max(0.0, min(1.0, t))
    a = [int(low[k:k + 2], 16) for k in (0, 2, 4)]
    b = [int(high[k:k + 2], 16) for k in (0, 2, 4)]
    return "".join(f"{round(a[k] + (b[k] - a[k]) * t):02X}" for k in range(3))


def bar_chart(slide, x, y, w, h, categories, series, colors, title=None, number_format="0.00", font_size=11,
              value_axis_title=None, max_value=None, min_value=None):
    data = CategoryChartData()
    data.categories = categories
    for name, values in series:
        data.add_series(name, values)
    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data)
    chart = frame.chart
    chart.font.size = Pt(font_size)
    chart.font.name = JP
    chart.font.color.rgb = rgb(BODY)
    chart.has_title = bool(title)
    if title:
        chart.chart_title.text_frame.text = title
        _style_runs(chart.chart_title.text_frame, 12, INK, True)
    chart.has_legend = len(series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(font_size)
    plot = chart.plots[0]
    plot.gap_width = 60
    plot.overlap = -10
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = number_format
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.size = Pt(font_size - 1)
    labels.font.color.rgb = rgb(INK)
    for s, color in zip(plot.series, colors):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = rgb(color)
    value_axis = chart.value_axis
    value_axis.has_major_gridlines = True
    value_axis.major_gridlines.format.line.color.rgb = rgb(LINE)
    value_axis.format.line.fill.background()
    value_axis.tick_labels.font.size = Pt(font_size - 1)
    value_axis.tick_labels.font.color.rgb = rgb(MUTED)
    if max_value is not None:
        value_axis.maximum_scale = max_value
    if min_value is not None:
        value_axis.minimum_scale = min_value
    if value_axis_title:
        value_axis.has_title = True
        value_axis.axis_title.text_frame.text = value_axis_title
        _style_runs(value_axis.axis_title.text_frame, font_size, MUTED, False)
    category_axis = chart.category_axis
    category_axis.tick_labels.font.size = Pt(font_size - 1)
    category_axis.tick_labels.font.color.rgb = rgb(INK)
    category_axis.format.line.color.rgb = rgb(LINE)
    category_axis.has_major_gridlines = False
    return frame


def doughnut(slide, x, y, w, h, categories, values, colors, hole=62, min_label_share=0.03):
    data = CategoryChartData()
    data.categories = categories
    data.add_series("内訳", values)
    frame = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(x), Inches(y), Inches(w), Inches(h), data)
    chart = frame.chart
    chart.has_legend = False
    chart.has_title = False
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.number_format = "0%"
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.show_percentage = True
    plot.data_labels.show_value = False
    plot.data_labels.font.size = Pt(12)
    plot.data_labels.font.bold = True
    plot.data_labels.font.color.rgb = rgb(WHITE)
    for index, point in enumerate(plot.series[0].points):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = rgb(colors[index])
        point.format.line.color.rgb = rgb(NAVY)
        if values[index] < min_label_share:
            # 小さい割合は輪の中に収まらないので、数字を出さない（凡例の文字に書く）
            label = point.data_label._get_or_add_dLbl()
            for child in list(label):
                if child.tag != qn("c:idx"):
                    label.remove(child)
            etree.SubElement(label, qn("c:delete")).set("val", "1")
    hole_size = chart._chartSpace.find(".//" + qn("c:holeSize"))
    if hole_size is None:
        doughnut_chart = chart._chartSpace.find(".//" + qn("c:doughnutChart"))
        hole_size = etree.SubElement(doughnut_chart, qn("c:holeSize"))
    hole_size.set("val", str(hole))
    return frame


def _style_runs(frame, size, color, bold):
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            _set_font(run, size, color, bold, False, JP)


def notes(slide, content):
    slide.notes_slide.notes_text_frame.text = content
